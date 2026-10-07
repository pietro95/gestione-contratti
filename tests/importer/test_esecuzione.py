"""F11.4b: import dell'esecuzione dal modello XLSX (SPEC §7). Dati di fantasia."""

import datetime
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse

from accounts.roles import EDITOR
from contratti import acquisti
from contratti import models as m
from contratti.importazioni import TIPI, modello_esecuzione, tipi_per
from contratti.importer.esecuzione import FOGLI, importa_esecuzione

from .. import factories as f

pytestmark = pytest.mark.django_db
D = datetime.date


def _file(tmp_path, fogli, nome="esecuzione.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for titolo, righe in fogli.items():
        ws = wb.create_sheet(titolo)
        ws.append(FOGLI[titolo])
        for r in righe:
            ws.append([r.get(c) for c in FOGLI[titolo]])
    percorso = tmp_path / nome
    wb.save(percorso)
    return percorso


@pytest.fixture
def contratto(make_user):
    make_user("rossi", role=EDITOR)
    it = f.cluster("IT")
    m.ContoCoge.objects.create(
        codice="600010", descrizione="Conto di prova", valido_dal=D(2020, 1, 1), cluster=it
    )
    c = f.contratto(
        fornitore=f.fornitore("Fornitore Prova A S.r.l."), numero_contratto="5800000016",
        cig="A000000001", data_inizio=D(2026, 3, 2), durata_mesi=1, una_tantum=True,
    )  # fmt: skip
    f.voce(c, importo_fisso=Decimal("10000.00"), massimale_variabile=Decimal("2000.00"))
    return c


FILE = {
    "RDA": [
        {"numero_rda": "5100000016", "cluster": "it", "descrizione": "Forniture di prova",
         "data_inserimento": "20/02/2026", "data_rilascio": "25/02/2026", "referente": "rossi",
         "numero_contratto": "5800000016", "posizione": "Hardware", "quantita": 2,
         "importo_unitario": "4.000,00", "conto": "600010", "centro_costo": "C0001"},
        {"numero_rda": "5100000016", "posizione": "Installazione", "importo_unitario": 2000,
         "di_cui_variabile": 2000},
        {"numero_rda": "4500000001", "cluster": "IT", "descrizione": "Numero non SAP",
         "posizione": "X", "importo_unitario": 1},
    ],
    "Fatture": [
        {"numero_contratto": "5800000016", "numero": "FA-1", "data": "15/03/2026",
         "importo_documento": 4000, "valore_base": 4000},
        {"numero_contratto": "A000000001", "numero": "FA-2", "data": "20/04/2026",
         "importo_documento": "5.000,00", "valore_base": 5000, "trasferita_il": "30/04/2026"},
        {"fornitore": "Fornitore Prova A", "numero": "FA-9", "data": "01/05/2026",
         "importo_documento": 10},
        {"numero_contratto": "5899999999", "numero": "FA-3", "data": "01/05/2026",
         "importo_documento": 1},
    ],
    "Documenti": [
        {"numero_contratto": "5800000016", "tipo": "DDT", "numero": "DDT-1", "data": "10/03/2026"},
        {"numero_contratto": "5800000016", "tipo": "ricezione", "data": "28/04/2026",
         "finale": "sì"},
        {"numero_contratto": "5800000016", "tipo": "ODA/RDC", "numero": "RDC-1",
         "data": "16/03/2026", "numero_fattura": "FA-1"},
        {"numero_contratto": "5800000016", "tipo": "fattura", "data": "01/01/2026"},
    ],
}  # fmt: skip


def test_import_completo_tc16(contratto, tmp_path):
    report = importa_esecuzione(_file(tmp_path, FILE), oggi=D(2026, 5, 1))
    c = report.conteggi
    assert (c["rda_create"], c["fatture_create"], c["documenti_creati"]) == (1, 3, 3)
    assert sorted(a.codice for a in report.anomalie) == [
        "CONTRATTO_NON_TROVATO", "NUMERO_RDA_NON_VALIDO", "TIPO_NON_VALIDO",
    ]  # fmt: skip
    rda = m.RichiestaAcquisto.objects.get(numero_rda="5100000016")
    assert rda.referente_utente.username == "rossi" and rda.cluster.codice == "IT"
    assert acquisti.importo_rda(rda) == Decimal("10000.00")
    assert list(rda.posizioni.values_list("di_cui_variabile", flat=True)) == [0, 2000]
    contratto.refresh_from_db()
    assert contratto.rda == rda
    # TC-16 al 30/04/2026: fatturato 9.000, residuo 1.000, fatture mancanti.
    riepilogo = acquisti.riepilogo_esecuzione(contratto, D(2026, 4, 30))
    assert riepilogo.esecuzione.fatturato == Decimal("9000.00")
    assert riepilogo.motivo_da_chiudere.value == "fatture_mancanti"
    fa2 = contratto.fatture.get(numero="FA-2")
    assert fa2.data_trasferimento == D(2026, 4, 30) and fa2.importo_documento == Decimal("5000")
    assert contratto.documenti_esecuzione.get(tipo="ODA_RDC").fattura.numero == "FA-1"
    assert contratto.documenti_esecuzione.get(tipo="RICEZIONE").finale
    # Senza contratto: la fattura entra per fornitore, da prendere in carico.
    assert m.Fattura.objects.get(numero="FA-9").contratto is None


def test_idempotenza_e_contratto_chiuso(contratto, tmp_path):
    percorso = _file(tmp_path, FILE)
    importa_esecuzione(percorso, oggi=D(2026, 5, 1))
    report = importa_esecuzione(percorso, oggi=D(2026, 5, 1))
    c = report.conteggi
    assert (c["rda_create"], c["fatture_create"], c["documenti_creati"]) == (0, 0, 0)
    assert (c["rda_gia_presenti"], c["fatture_gia_presenti"], c["documenti_gia_presenti"]) == (
        1, 3, 3,
    )  # fmt: skip
    acquisti.chiudi(contratto, None, D(2026, 5, 10))
    nuova = {"Fatture": [{"numero_contratto": "5800000016", "numero": "FA-4",
                          "data": "15/05/2026", "importo_documento": 1}]}  # fmt: skip
    report = importa_esecuzione(_file(tmp_path, nuova, "chiuso.xlsx"))
    assert [a.codice for a in report.anomalie] == ["FATTURA_DOPO_CHIUSURA"]


def test_dry_run(contratto, tmp_path):
    report = importa_esecuzione(_file(tmp_path, FILE), dry_run=True)
    assert report.conteggi["fatture_create"] == 3
    assert not m.Fattura.objects.exists()


def test_tipi_solo_admin_e_modello(client, make_user):
    editor = make_user("editor_x", role=EDITOR)
    assert {"esecuzione", "riepilogo"}.isdisjoint(t.codice for t in tipi_per(editor))
    wb = modello_esecuzione()
    assert wb.sheetnames == ["RDA", "Documenti", "Fatture", "Istruzioni"]
    assert TIPI["riepilogo"].solo_admin and TIPI["esecuzione"].solo_admin
    client.force_login(editor)
    url = reverse("contratti:importazione_modello", args=["esecuzione"])
    assert client.get(url).status_code == 404
