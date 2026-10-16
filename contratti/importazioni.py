"""Import dall'interfaccia (F9.1): caricamento, anteprima senza scritture, conferma.

Gli import sono gli stessi della riga di comando (stesso codice): Excel dei contratti, date
di emissione, piano dei conti, quadratura con l'Excel. Il file resta nel registro
(``Importazione``) con anteprima ed esito.
"""

import datetime
import hashlib
import io
import logging
import re
import tempfile
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from accounts.roles import ADMIN, WRITE_ROLES, role_of
from contratti import models as m
from contratti import permessi
from contratti.importer import modello_struttura as modello_struttura_rev00
from contratti.importer.carica import importa as importa_excel
from contratti.importer.conti import importa_conti
from contratti.importer.date_emissione import importa_date
from contratti.importer.esecuzione import FOGLI as FOGLI_ESECUZIONE
from contratti.importer.esecuzione import OBBLIGATORIE as OBBLIGATORIE_ESECUZIONE
from contratti.importer.esecuzione import importa_esecuzione
from contratti.importer.fatture import importa_fatture
from contratti.importer.pregresso import importa_pregresso
from contratti.importer.report import Report
from contratti.importer.riepilogo import importa_riepilogo
from contratti.importer.struttura import importa_struttura
from contratti.importer.tabelle import FileNonValido
from contratti.importer.tabelle_sap import importa_tabelle_sap

log = logging.getLogger(__name__)

DIMENSIONE_MASSIMA = 10 * 1024 * 1024  # 10 MB: i file reali sono di poche centinaia di righe
# Un XLSX è uno zip: oltre questa dimensione decompressa (o questo rapporto) è sospetto, e
# leggerlo potrebbe esaurire la memoria del container (zip bomb).
DECOMPRESSO_MASSIMO = 200 * 1024 * 1024
RAPPORTO_MASSIMO = 200
VALIDITA_ANTEPRIMA = datetime.timedelta(hours=24)
S = m.Importazione.Stato


class ImportNonAmmesso(Exception):
    pass


@dataclass(frozen=True)
class Tipo:
    codice: str
    titolo: str
    descrizione: str
    estensioni: tuple[str, ...]
    esegui: Callable  # (percorso, dry_run, opzioni) -> dict serializzato
    etichette: dict[str, str]  # conteggi da mostrare, nell'ordine
    solo_admin: bool = False
    con_conferma: bool = True
    chiede_area: bool = False
    modello: Callable[[], Workbook] | None = None
    note: list[str] = field(default_factory=list)
    oggetti_modello: str = ""  # nel nome del file: GestioneContratti_ModelloImport_<oggetti>.xlsx


def serializza(report: Report) -> dict:
    return {
        "conteggi": dict(report.conteggi),
        "anomalie": [
            [a.foglio, a.riga, a.codice, a.campo, a.valore, a.messaggio] for a in report.anomalie
        ],
        "mappature": [[v, n, k] for (v, n), k in sorted(report.mappature_conti.items())],
    }


# --- Esecuzione dei singoli tipi ------------------------------------------------------


def _contratti(percorso, dry_run, opzioni):
    return serializza(importa_excel(percorso, area=opzioni["area"], dry_run=dry_run)), None


def _date(percorso, dry_run, opzioni):
    ammessi = opzioni.get("cluster_ammessi")  # None: tutti i cluster
    contratti = None
    if ammessi is not None:
        contratti = m.Contratto.objects.filter(fabbisogno__cluster_id__in=ammessi)
    return serializza(importa_date(percorso, dry_run=dry_run, contratti=contratti)), None


def _pregresso(percorso, dry_run, opzioni):
    ammessi = opzioni.get("cluster_ammessi")  # None: tutti i cluster
    contratti = None
    if ammessi is not None:
        contratti = m.Contratto.objects.filter(fabbisogno__cluster_id__in=ammessi)
    return serializza(importa_pregresso(percorso, dry_run=dry_run, contratti=contratti)), None


def _conti(percorso, dry_run, opzioni):
    return serializza(importa_conti(percorso, dry_run=dry_run)), None


def _struttura(percorso, dry_run, opzioni):
    return serializza(importa_struttura(percorso, dry_run=dry_run)), None


def _tabelle_sap(percorso, dry_run, opzioni):
    return serializza(importa_tabelle_sap(percorso, dry_run=dry_run)), None


def _fatture(percorso, dry_run, opzioni):
    return serializza(importa_fatture(percorso, dry_run=dry_run)), None


def _esecuzione(percorso, dry_run, opzioni):
    return serializza(importa_esecuzione(percorso, dry_run=dry_run)), None


def _riepilogo(percorso, dry_run, opzioni):
    return serializza(importa_riepilogo(percorso, dry_run=dry_run)), None


def _quadratura(percorso, dry_run, opzioni):
    from reports import xlsx
    from reports.export import quadratura_xlsx
    from reports.quadratura import CAUSE, quadra

    q = quadra(percorso, area=opzioni["area"])
    dati = {
        "conteggi": {"righe_confrontate": len(q.righe), "righe_escluse": len(q.escluse)},
        "anomalie": [["Dashboard", n, "RIGA_ESCLUSA", "", "", motivo] for n, motivo in q.escluse]
        + [["", None, c, "", k, CAUSE[c]] for c, k in q.conteggio_cause().items()],
        "riepilogo": [
            {k: (str(v) if v is not None else None) for k, v in r.items()} for r in q.riepilogo()
        ],
    }
    return dati, xlsx.salva(quadratura_xlsx(q))


# --- Modelli da scaricare -------------------------------------------------------------

GIALLO = PatternFill("solid", fgColor="FFF2CC")


def _modello(fogli: list[tuple[str, list[str], set[str]]], istruzioni: list[str]) -> Workbook:
    """Fogli dati con l'intestazione (gialla = obbligatoria), poi il foglio Istruzioni."""
    wb = Workbook()
    wb.remove(wb.active)
    for titolo, colonne, obbligatorie in fogli:
        ws = wb.create_sheet(titolo)
        for j, nome in enumerate(colonne, start=1):
            c = ws.cell(1, j, nome)
            c.font = Font(name="Arial", bold=True)
            if nome in obbligatorie:
                c.fill = GIALLO
            ws.column_dimensions[c.column_letter].width = max(14, len(nome) + 4)
        ws.freeze_panes = "A2"
    ws = wb.create_sheet("Istruzioni")
    ws.column_dimensions["A"].width = 110
    for i, testo in enumerate(istruzioni, start=1):
        ws.cell(i, 1, testo).font = Font(name="Arial", bold=i == 1)
    return wb


def modello_struttura() -> Workbook:
    """Il modello approvato da Pietro (Rev00), con istruzioni e colonna di controllo."""
    return modello_struttura_rev00.crea()


def modello_date() -> Workbook:
    return _modello(
        [("Date", ["numero_contratto", "data_emissione"], {"numero_contratto", "data_emissione"})],
        [
            "Date di emissione dei contratti (estrazione SAP)",
            "Una riga per contratto: numero di contratto SAP e data di emissione (gg/mm/aaaa).",
            "Si compilano solo le date vuote; le date già registrate e diverse vanno nel report.",
            "Esempio (fittizio): 5800000000 | 15/12/2025",
        ],
    )


def modello_pregresso() -> Workbook:
    return _modello(
        [
            (
                "Fatturato",
                ["numero_contratto", "id_contratto", "anno", "fatturato", "nota"],
                {"anno", "fatturato"},
            )
        ],
        [
            "Fatturato pregresso dei contratti (totale per anno)",
            "Una riga per contratto e anno. Il contratto si indica con il numero di contratto "
            "SAP oppure, se il numero manca o è su più contratti, con l'id (l'ultimo numero "
            "nell'indirizzo della pagina del contratto). Serve una delle due colonne.",
            "fatturato: totale dell'anno, imponibile (IVA esclusa), come il valore base delle "
            "fatture. Conta nel fatturato, nel residuo e nell'avanzamento del contratto.",
            "Rieseguire il file non duplica: lo stesso anno con lo stesso importo non cambia, con "
            "un importo diverso corregge il totale. Le righe non applicabili vanno nel report.",
            "Esempio (fittizio): 5800000000 | | 2024 | 12000,00 | da Excel di cluster",
        ],
    )


def modello_conti() -> Workbook:
    return _modello(
        [
            (
                "Conti",
                ["codice", "descrizione", "valido_dal", "valido_al", "cluster"],
                {"codice", "descrizione", "valido_dal"},
            )
        ],
        [
            "Piano dei conti",
            "La coppia codice + valido_dal identifica il conto; una riga già presente aggiorna "
            "descrizione, fine validità e cluster. Nessuna cancellazione.",
            "cluster: codice del cluster del conto (es. IT); vuoto = invariato. Ogni conto "
            "appartiene a un cluster: per caricare tutto insieme usa l'import della struttura.",
            "Date gg/mm/aaaa. Un errore su una riga annulla tutto il file.",
            "Esempio (fittizio): 60000000 | Canoni software | 01/01/2026 | | IT",
        ],
    )


def modello_tabelle_sap() -> Workbook:
    return _modello(
        [
            ("Richiedenti", ["codice", "descrizione", "predefinito"], {"codice"}),
            ("Gruppi merci", ["codice", "descrizione", "predefinito"], {"codice"}),
            ("Gruppi acquisti", ["codice", "nome", "predefinito"], {"codice"}),
            ("Centri di costo", ["codice", "descrizione", "predefinito"], {"codice"}),
        ],
        [
            "Tabelle SAP del processo d'acquisto",
            "Un foglio per tabella; si può caricare anche un foglio solo. Il codice identifica "
            "la riga: una riga già presente aggiorna descrizione (o nome). Nessuna "
            "cancellazione: un codice tolto dal file resta, si disattiva dall'amministrazione.",
            "predefinito: «sì» su una sola riga per foglio (es. il richiedente, il gruppo "
            "acquisti e il centro di costo dell'ufficio); diventa l'unico predefinito e si "
            "propone su RDA, voci e posizioni nuove. Vuoto = no.",
            "Centri di costo: dopo il caricamento il report elenca i codici già usati da voci e "
            "posizioni che la tabella non ha (CENTRO_COSTO_NON_IN_TABELLA).",
            "Un errore (due predefiniti, codici ripetuti) annulla il file.",
            "Esempio (fittizio): 999901 | Ufficio di prova | sì",
        ],
    )


def modello_fatture() -> Workbook:
    return _modello(
        [
            (
                "Fatture",
                ["Fornitore", "Numero Documento", "Data", "Importo SAP"],
                {"Fornitore", "Numero Documento", "Data", "Importo SAP"},
            )
        ],
        [
            "Fatture dall'estrazione SAP settimanale",
            "Si legge il primo foglio; l'intestazione è la riga con «Fornitore» e «Numero "
            "Documento» (sopra può esserci la riga di revisione). Le altre colonne si ignorano.",
            "Fornitore: ragione sociale o nome breve come in anagrafica. Data gg.mm.aaaa, "
            "gg.mm.aa o gg/mm/aaaa. Importo SAP: importo del documento.",
            "Le fatture entrano da prendere in carico, senza contratto; una fattura già "
            "registrata (stesso fornitore e numero) non si duplica e non si modifica.",
            "Esempio (fittizio): Fornitore Prova A S.r.l. | FA-1/2026 | 04.09.2026 | 1220",
        ],
    )


def modello_esecuzione() -> Workbook:
    return _modello(
        [
            (nome, colonne, OBBLIGATORIE_ESECUZIONE[nome])
            for nome, colonne in FOGLI_ESECUZIONE.items()
        ],
        [
            "Esecuzione: RDA, documenti di esecuzione e fatture letti in SAP",
            "Fogli facoltativi. RDA: una riga per posizione, le righe con lo stesso numero RDA "
            "sono la stessa RDA; cluster con il codice (es. IT); referente con il nome utente.",
            "numero_contratto: numero di contratto SAP o CIG del contratto già registrato.",
            "Documenti: tipo ODA/RDC, DDT, MIGO, MAP, attestazione, ricezione; finale «sì» solo "
            "per attestazione e ricezione; numero_fattura per legarlo a una fattura.",
            "Fatture: senza numero_contratto serve il fornitore (ragione sociale o nome breve); "
            "trasferita_il = data del trasferimento per il pagamento.",
            "Date gg/mm/aaaa. Rieseguire l'import con lo stesso file non crea duplicati.",
        ],
    )


def modello_contratti() -> Workbook:
    return _modello(
        [
            (
                "Contratti",
                [
                    "Contratto",
                    "Cig",
                    "Nuovo codice conto da 01/01/26",
                    "Codice conto",
                    "Descrizione",
                    "Fornitore",
                    "Codice fornitore",
                    "Corrispettivo Minimo",
                    "Corrispettivo Extra",
                    "Corrispettivo Massimo",
                    "Data di attivazione",
                    "Mesi contratto",
                    "Data di scadenza",
                    "Stato",
                    "Note",
                ],  # fmt: skip
                {"Contratto", "Descrizione", "Fornitore", "Corrispettivo Minimo"},
            ),
            ("Fornitori", ["Fornitore", "Codice"], {"Fornitore"}),
        ],
        [
            "Excel dei contratti (stessa struttura dei file Contratti_IT / Contratti_TLC)",
            "Si leggono i fogli il cui nome inizia con «Contratti» (i fogli «Contratti_Finanziati» "
            "e «Contratti Dissalatori» sono spese su WBS: colonna «Codice WBS»).",
            "Ogni riga entra in un fabbisogno provvisorio; raggruppamenti e catene dei rinnovi "
            "diventano proposte da confermare in «Proposte import».",
            "Rieseguire l'import con lo stesso file non crea duplicati e non sovrascrive le "
            "modifiche fatte dopo.",
        ],
    )


TIPI: dict[str, Tipo] = {
    t.codice: t
    for t in (
        Tipo(
            "contratti",
            "Excel dei contratti",
            "Contratti, voci e fornitori da un Excel con la struttura dei file di partenza.",
            (".xlsx", ".xlsm"),
            _contratti,
            {
                "righe_lette": "Righe lette",
                "contratti_creati": "Contratti creati",
                "contratti_gia_presenti": "Già presenti (scartati)",
                "voci_create": "Voci create",
                "fornitori_creati": "Fornitori creati",
                "proposte_raggruppamento": "Proposte di raggruppamento",
                "proposte_predecessore": "Proposte di catena",
            },
            chiede_area=True,
            modello=modello_contratti,
            oggetti_modello="Contratti",
        ),
        Tipo(
            "date_emissione",
            "Date di emissione",
            "Data di emissione dei contratti già registrati, da un'estrazione SAP.",
            (".xlsx", ".csv"),
            _date,
            {
                "compilati": "Date compilate",
                "invariati": "Già uguali",
                "discordanti": "Discordanti (non modificate)",
                "non_trovati": "Numeri non trovati",
                "fuori_cluster": "Di altri cluster (non modificati)",
            },
            modello=modello_date,
            oggetti_modello="DateEmissione",
        ),
        Tipo(
            "pregresso",
            "Fatturato pregresso",
            "Totale fatturato per anno dei contratti in corso o chiusi, da un foglio con numero "
            "di contratto (o id), anno e importo imponibile.",
            (".xlsx", ".csv"),
            _pregresso,
            {
                "creati": "Totali creati",
                "corretti": "Totali corretti",
                "invariati": "Già uguali",
                "scartati": "Righe scartate",
            },
            modello=modello_pregresso,
            oggetti_modello="FatturatoPregresso",
        ),
        Tipo(
            "conti",
            "Piano dei conti",
            "Conti COGE con validità temporale. Tabella di configurazione: solo Admin.",
            (".xlsx", ".csv"),
            _conti,
            {"creati": "Conti creati", "aggiornati": "Aggiornati", "invariati": "Invariati"},
            solo_admin=True,
            modello=modello_conti,
            oggetti_modello="PianoConti",
        ),
        Tipo(
            "struttura",
            "Struttura ICTDINN",
            "Cluster, responsabili, piano dei conti e WBS. Configurazione: solo Admin.",
            (".xlsx",),
            _struttura,
            {
                "cluster_creati": "Cluster creati",
                "cluster_aggiornati": "Cluster aggiornati",
                "assegnazioni_create": "Responsabili aggiunti",
                "assegnazioni_tolte": "Responsabili tolti",
                "conti_creati": "Conti creati",
                "conti_aggiornati": "Conti aggiornati",
                "wbs_create": "WBS create",
                "wbs_aggiornate": "WBS aggiornate",
                "voci_collegate": "Voci collegate alla WBS",
            },
            solo_admin=True,
            modello=modello_struttura,
            oggetti_modello="Struttura",
            note=[
                "Il foglio Responsabili, se compilato, è l'elenco completo: chi non c'è perde "
                "le assegnazioni. Conti e WBS non si cancellano mai.",
            ],
        ),
        Tipo(
            "tabelle_sap",
            "Tabelle SAP (RDA)",
            "Richiedenti SAP, gruppi merci, gruppi acquisti e centri di costo: il "
            "modello, oppure gli elenchi di SAP così come sono (gruppi merci; codici di "
            "rilascio con unità, richiedente e gruppo acquisti). Configurazione: solo Admin.",
            (".xlsx",),
            _tabelle_sap,
            {
                "richiedenti_creati": "Richiedenti creati",
                "richiedenti_aggiornati": "Richiedenti aggiornati",
                "gruppi_merci_creati": "Gruppi merci creati",
                "gruppi_merci_aggiornati": "Gruppi merci aggiornati",
                "gruppi_acquisti_creati": "Gruppi acquisti creati",
                "gruppi_acquisti_aggiornati": "Gruppi acquisti aggiornati",
                "richiedenti_collegati": "Richiedenti legati al gruppo acquisti",
                "centri_costo_creati": "Centri di costo creati",
                "centri_costo_aggiornati": "Centri di costo aggiornati",
                "centri_costo_non_in_tabella": "Centri di costo usati e assenti dalla tabella",
            },
            solo_admin=True,
            modello=modello_tabelle_sap,
            oggetti_modello="TabelleSAP",
        ),
        Tipo(
            "fatture",
            "Fatture (estrazione SAP)",
            "Import settimanale delle fatture: entrano da prendere in carico, con il contratto "
            "e il referente proposti.",
            (".xlsx", ".xlsm"),
            _fatture,
            {
                "righe_lette": "Righe lette",
                "create": "Fatture create",
                "gia_presenti": "Già registrate",
                "non_importate": "Non importate",
                "fornitori_creati": "Fornitori creati",
                "con_contratto_proposto": "Con contratto proposto",
            },
            modello=modello_fatture,
            oggetti_modello="Fatture",
            note=[
                "Dopo la conferma le fatture sono in «Fatture da prendere in carico», visibili a "
                "tutti.",
            ],
        ),
        Tipo(
            "esecuzione",
            "Esecuzione (modello XLSX)",
            "RDA con posizioni, documenti di esecuzione e fatture letti in SAP, riportati nel "
            "modello. Solo Admin.",
            (".xlsx",),
            _esecuzione,
            {
                "rda_create": "RDA create",
                "rda_gia_presenti": "RDA già registrate",
                "documenti_creati": "Documenti creati",
                "documenti_gia_presenti": "Documenti già registrati",
                "fatture_create": "Fatture create",
                "fatture_gia_presenti": "Fatture già registrate",
            },
            solo_admin=True,
            modello=modello_esecuzione,
            oggetti_modello="Esecuzione",
        ),
        Tipo(
            "riepilogo",
            "Riepilogo RDA Ordini (una tantum)",
            "Il file «Riepilogo RDA Ordini»: RDA ICTDINN dal 01/04/2025 con contratti, fatture, "
            "DDT, attestazioni e numeri SAP. Solo Admin.",
            (".xlsx",),
            _riepilogo,
            {
                "rda_lette": "RDA lette",
                "rda_create": "RDA create",
                "rda_gia_presenti": "RDA già registrate",
                "fuori_perimetro": "Fuori perimetro",
                "contratti_collegati": "Contratti collegati",
                "contratti_creati": "Contratti creati",
                "fornitori_creati": "Fornitori creati",
                "fatture_create": "Fatture create",
                "documenti_creati": "Documenti creati",
                "contratti_chiusi": "Contratti chiusi",
            },
            solo_admin=True,
            note=[
                "I contratti non registrati si creano nel fabbisogno «Acquisti una tantum "
                "<cluster> (dal Riepilogo)», da ricollocare; il referente delle RDA resta vuoto.",
            ],
        ),
        Tipo(
            "quadratura",
            "Quadratura con l'Excel",
            "Confronta la competenza del registro con il Dashboard dell'Excel. Non scrive nulla.",
            (".xlsx", ".xlsm"),
            _quadratura,
            {"righe_confrontate": "Righe confrontate", "righe_escluse": "Righe escluse"},
            con_conferma=False,
            chiede_area=True,
        ),
    )
}


# --- Permessi e operazioni --------------------------------------------------------------


def tipi_per(utente) -> list[Tipo]:
    ruolo = role_of(utente)
    if ruolo not in WRITE_ROLES:
        return []
    return [t for t in TIPI.values() if ruolo == ADMIN or not t.solo_admin]


def cluster_importabili(tipo: Tipo, utente) -> list[tuple[str, str]]:
    """Cluster scelti all'import: per i contratti solo quelli dove l'utente modifica (F9.2)."""
    scelte = m.scelte_cluster()
    if tipo.codice != "contratti" or utente is None:
        return scelte
    ammessi = permessi.cluster_modificabili(utente)
    if ammessi is None:
        return scelte
    codici = set(m.Cluster.objects.filter(pk__in=ammessi).values_list("codice", flat=True))
    return [(c, n) for c, n in scelte if c in codici]


def puo_usare(utente, tipo: Tipo) -> bool:
    return tipo in tipi_per(utente)


def puo_confermare(utente, imp: m.Importazione) -> bool:
    tipo = TIPI.get(imp.tipo)
    if tipo is None or not puo_usare(utente, tipo):
        return False
    return role_of(utente) == ADMIN or imp.utente_id == utente.pk


def _nome_sicuro(nome: str) -> str:
    base = Path(nome).name
    return re.sub(r"[^\w.\- ]", "_", base)[:120] or "file"


def _esegui(imp: m.Importazione, *, dry_run: bool):
    """Scrive il file in una cartella temporanea con il suo nome e lancia l'import."""
    tipo = TIPI[imp.tipo]
    with tempfile.TemporaryDirectory(prefix="import-") as cartella:
        percorso = Path(cartella) / _nome_sicuro(imp.nome_file)
        percorso.write_bytes(bytes(imp.contenuto))
        return tipo.esegui(percorso, dry_run, imp.opzioni)


def controlla_xlsx(contenuto: bytes) -> None:
    """Rifiuta gli XLSX che non sono zip validi o che si espandono troppo."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenuto)) as z:
            decompresso = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile as exc:
        raise ImportNonAmmesso("Il file non è un Excel valido (.xlsx).") from exc
    if decompresso > DECOMPRESSO_MASSIMO or decompresso > RAPPORTO_MASSIMO * max(len(contenuto), 1):
        raise ImportNonAmmesso("Il file Excel è troppo grande una volta decompresso.")


def carica(tipo: Tipo, file, opzioni: dict, utente) -> m.Importazione:
    """Registra il file e ne calcola l'anteprima (o il report, per i tipi senza conferma)."""
    if not puo_usare(utente, tipo):
        raise ImportNonAmmesso("Questo import non è disponibile per il tuo ruolo.")
    if tipo.codice == "contratti" and opzioni.get("area") not in dict(
        cluster_importabili(tipo, utente)
    ):
        raise ImportNonAmmesso("Puoi importare contratti solo nei tuoi cluster.")
    if tipo.codice in ("date_emissione", "pregresso"):
        ammessi = permessi.cluster_modificabili(utente)
        opzioni = {**opzioni, "cluster_ammessi": None if ammessi is None else sorted(ammessi)}
    nome = file.name or "file"
    if Path(nome).suffix.lower() not in tipo.estensioni:
        raise ImportNonAmmesso(f"Formato non ammesso: servono file {', '.join(tipo.estensioni)}.")
    if file.size > DIMENSIONE_MASSIMA:
        raise ImportNonAmmesso("Il file supera i 10 MB.")
    contenuto = file.read()
    if Path(nome).suffix.lower() in (".xlsx", ".xlsm"):
        controlla_xlsx(contenuto)
    imp = m.Importazione.objects.create(
        tipo=tipo.codice,
        nome_file=nome[:255],
        contenuto=contenuto,
        sha256=hashlib.sha256(contenuto).hexdigest(),
        dimensione=len(contenuto),
        opzioni=opzioni,
        utente=utente,
    )
    try:
        dati, risultato = _esegui(imp, dry_run=True)
    except FileNonValido as exc:
        imp.stato, imp.messaggio = S.ERRORE, str(exc)
    except Exception as exc:  # file inatteso: si registra, non si interrompe la pagina
        log.exception("Anteprima dell'import %s non riuscita", imp.pk)
        imp.stato, imp.messaggio = S.ERRORE, f"Il file non si può leggere: {exc}"
    else:
        imp.anteprima, imp.risultato = dati, risultato
        if not tipo.con_conferma:
            imp.stato = S.COMPLETATA
    imp.save()
    return imp


def gia_importato(imp: m.Importazione) -> m.Importazione | None:
    """Lo stesso file (stesso contenuto) già confermato per lo stesso tipo."""
    return (
        m.Importazione.objects.filter(tipo=imp.tipo, sha256=imp.sha256, stato=S.CONFERMATA)
        .exclude(pk=imp.pk)
        .first()
    )


def scaduta(imp: m.Importazione, adesso=None) -> bool:
    adesso = adesso or timezone.now()
    return imp.stato == S.ANTEPRIMA and adesso - imp.caricata_il > VALIDITA_ANTEPRIMA


def conferma(pk: int, utente) -> m.Importazione:
    """Applica il file dell'anteprima. Una sola volta: la riga è bloccata durante l'import.

    Gli esiti negativi (scaduta, errore) si salvano e solo dopo, fuori dalla transazione,
    si segnalano: un'eccezione dentro la transazione annullerebbe anche lo stato.
    """
    errore = None
    with transaction.atomic():
        imp = m.Importazione.objects.select_for_update().get(pk=pk)
        if not puo_confermare(utente, imp):
            raise ImportNonAmmesso("Può confermare solo chi ha caricato il file, o un Admin.")
        if imp.stato != S.ANTEPRIMA:
            raise ImportNonAmmesso("Questo import non è più in anteprima.")
        if scaduta(imp):
            imp.stato = S.SCADUTA
            imp.save(update_fields=["stato"])
            errore = "Anteprima scaduta (oltre 24 ore): carica di nuovo il file."
        elif hashlib.sha256(bytes(imp.contenuto)).hexdigest() != imp.sha256:
            errore = "Il file registrato non corrisponde: carica di nuovo il file."
        else:
            try:
                # L'import ha la sua transazione: se fallisce si annulla solo quella.
                dati, risultato = _esegui(imp, dry_run=False)
            except FileNonValido as exc:
                errore = str(exc)
            except Exception as exc:
                log.exception("Import %s non riuscito", imp.pk)
                errore = f"Import non riuscito, nessun dato scritto: {exc}"
            if errore:
                imp.stato, imp.messaggio = S.ERRORE, errore
            else:
                imp.esito, imp.risultato = dati, risultato or imp.risultato
                imp.stato = S.CONFERMATA
            imp.confermata_il = timezone.now()
            imp.confermata_da = utente
            imp.save()
    if errore:
        raise ImportNonAmmesso(errore)
    return imp


def annulla(pk: int, utente) -> m.Importazione:
    with transaction.atomic():
        imp = m.Importazione.objects.select_for_update().get(pk=pk)
        if not puo_confermare(utente, imp):
            raise ImportNonAmmesso("Può annullare solo chi ha caricato il file, o un Admin.")
        if imp.stato != S.ANTEPRIMA:
            raise ImportNonAmmesso("Questo import non è più in anteprima.")
        imp.stato = S.ANNULLATA
        imp.save(update_fields=["stato"])
    return imp


def report_xlsx(imp: m.Importazione, dati: dict) -> Workbook:
    """Conteggi e anomalie in XLSX, per l'anteprima o per l'esito."""
    from reports import xlsx

    tipo = TIPI[imp.tipo]
    wb = xlsx.nuovo_libro()
    intestazione = [
        f"{tipo.titolo} — {imp.nome_file}",
        f"Caricato da {imp.utente} il {timezone.localtime(imp.caricata_il):%d/%m/%Y %H:%M}",
    ]
    conteggi = dati.get("conteggi", {})
    righe = [[tipo.etichette.get(k, k), v] for k, v in conteggi.items()]
    xlsx.foglio(wb, "Riepilogo", intestazione, ["Voce", "Numero"], righe)
    colonne = ["Foglio", "Riga", "Codice", "Campo", "Valore", "Messaggio"]
    xlsx.foglio(wb, "Anomalie", intestazione, colonne, dati.get("anomalie", []))
    if dati.get("mappature"):
        xlsx.foglio(
            wb, "Mappature conti", intestazione, ["Conto vecchio", "Conto nuovo", "Righe"],
            dati["mappature"],
        )  # fmt: skip
    return wb
