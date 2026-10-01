"""Scrittura dei report in XLSX: numeri come numeri, date come date, intestazione con i filtri."""

import datetime
import io
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

FORMATO_IMPORTO = "#,##0.00"
FORMATO_DATA = "DD/MM/YYYY"
TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def nuovo_libro() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def foglio(
    wb: Workbook,
    titolo: str,
    intestazione: list[str],
    colonne: list[str],
    righe: list[list],
):
    """Un foglio: righe d'intestazione (titolo, filtri), poi tabella con le colonne."""
    ws = wb.create_sheet(titolo[:31])
    for i, testo in enumerate(intestazione, start=1):
        ws.cell(row=i, column=1, value=testo).font = Font(bold=i == 1)
    prima = len(intestazione) + 2
    for j, nome in enumerate(colonne, start=1):
        ws.cell(row=prima, column=j, value=nome).font = Font(bold=True)
    for i, riga in enumerate(righe, start=prima + 1):
        for j, valore in enumerate(riga, start=1):
            cella = ws.cell(row=i, column=j, value=valore)
            if isinstance(valore, Decimal):
                cella.number_format = FORMATO_IMPORTO
            elif isinstance(valore, datetime.date):
                cella.number_format = FORMATO_DATA
    for j, nome in enumerate(colonne, start=1):
        larghezza = max([len(str(nome))] + [len(_testo(r[j - 1])) for r in righe[:200]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(larghezza + 2, 10), 60)
    ws.freeze_panes = ws.cell(row=prima + 1, column=1)
    return ws


def _testo(valore) -> str:
    if isinstance(valore, Decimal):
        return f"{valore:,.2f}"
    if isinstance(valore, datetime.date):
        return "00/00/0000"
    return "" if valore is None else str(valore)


def salva(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def risposta(wb: Workbook, nome_file: str) -> HttpResponse:
    response = HttpResponse(salva(wb), content_type=TIPO_XLSX)
    response["Content-Disposition"] = f'attachment; filename="{nome_file}"'
    return response
