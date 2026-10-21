# ruff: noqa: E501 — righe di una tabella, come nel file reale
"""Riepilogo del budget in miniatura, con la stessa struttura del file reale (celle unite,
«(non utile)», «DA EVITARE», conti multipli in una cella) e solo valori fittizi."""

import openpyxl

INTESTAZIONE = [
    "Cluster", "SottoCluster", "Owner", "Tipo costo", None, "Codice", "Descrizione",
    "Codice \n(I livello)", "Descrizione \n(I livello)", "Codice \n(II livello)",
    "Descrizione \n(II livello)", "Codice \n(III livello)", "Descrizione \n(III livello)",
]  # fmt: skip


def crea(percorso, *, anno=2026):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget prova"
    ws["D2"] = "RIEPILOGO REPORTING (prova)"
    ws["A4"], ws["E4"], ws["F4"], ws["H4"] = (
        "Clusterizzazione budget (prova)", "CDC \nCentro di costo", "VDC / Co.Ge.", "WBS",
    )  # fmt: skip
    ws.merge_cells("F4:G4")
    ws.merge_cells("H4:M4")
    for i, v in enumerate(INTESTAZIONE, start=1):
        ws.cell(5, i, v)
    righe = [
        # cluster, sotto, owner, natura, cdc, conto, descr. conto, wbs I, desc I, wbs II, desc II,
        # wbs III, desc III
        ("ALFA", "SERVIZI RETE", "Mario Rossi", "CAPEX", 10520, 600001, "SERVIZI INTERNET",
         "VJ/00022\n(prima era VJ/00011)", "Ricerca", "VJ/00022/001\n(prima era VJ/00011/001)",
         "Ricerca e sviluppo", "--", "--"),
        (None, None, None, "OPEX", 10520, 600002, "TELEFONIA FISSA", "--", "--", "--", "--", "--", "--"),
        (None, None, None, None, None, 600001, "SERVIZI INTERNET", "--", "--", "--", "--", "--", "--"),
        (None, "ELABORAZIONE", "Mario Rossi", "CAPEX", 10520, 600010, " ACQUISTO MATERIALE",
         "VJ/00007", "Attrezzature", "VJ/00007/001", "Macchine d'ufficio", "--", "--"),
        (None, None, None, None, None, None, None, "VJ/00024\nDA EVITARE", "Noleggio",
         "VJ/00024/001\nDA EVITARE", "Uso macchine da terzi", "--", "--"),
        (None, None, None, "OPEX", 10520, 600010, " ACQUISTO MATERIALE", "--", "--", "--", "--", "--", "--"),
        ("BETA", "SICUREZZA FISICA\n(videosorveglianza, antintrusione)", "Anna Bianchi", "CAPEX", 10520,
         600020, "MONTAGGI", "VJ/00007", "Attrezzature", "VJ/00007/003", "Imp. Sicurezza fisica", "--", "--"),
        (None, "SICUREZZA INTEGRATA (intelligence)", "Anna Bianchi", "OPEX", 10520, "(da definire)",
         "(da definire)", "--", "--", "--", "--", "--", "--"),
        (None, "SICUREZZA \n(Holding)", None, "CDS Holding Overhead", None, None, None, None, None,
         None, None, None, None),
        (None, "IMPIANTI ELETTRICI\n(compresi hardware OT)", "Anna Bianchi", "CAPEX", "10520 (nuovo 21400)",
         "Per Forniture: 600031", "ACQUISTI MATERIALI", "SI.98.MS001", "MANUTENZIONE IN CORSO",
         "(non utile)", "(non utile)", "SI.98.MS001.01.9900", "(non utile)"),
        (None, None, None, None, None, "Per Servizi: 600032", "ALTRI SERVIZI", None, None, None,
         None, None, None),
        ("GAMMA", "INNOVAZIONE\n(PoC, sperimentazioni)", "Luca Verdi", "CAPEX", 10520,
         "Per Forniture: 600040\n\nPer Lavori: 600041",
         "Per Forniture: 600040 - MATERIALI DI CONSUMO\n\nPer Lavori: 600041 - SCAVI",
         "SI.98.MS002", "MS INNOVAZIONE", "(non utile)", "(non utile)", "SI.98.MS002.01.9900", "(non utile)"),
        (None, None, None, None, None, None, None, None, None, None, None, None, None),
        (None, None, None, "OPEX", "--", "--", "--", "--", "--", "--", "--", "--", "--"),
    ]  # fmt: skip
    for i, riga in enumerate(righe, start=6):
        for j, v in enumerate(riga, start=1):
            if v is not None:
                ws.cell(i, j, v)
    for area in (
        "A6:A11", "B6:B8", "C6:C8", "D7:D8", "B9:B11", "C9:C11", "D9:D10",
        "A12:A16", "B15:B16", "C15:C16", "D15:D16", "D14:M14",
        "A17:A19", "B17:B19", "C17:C19", "D17:D18",
        "F17:F18", "G17:G18", "H17:H18", "I17:I18", "L17:L18",
    ):  # fmt: skip
        ws.merge_cells(area)
    ws.cell(21, 11, f"TOTALE BUDGET {anno}")
    ws.cell(22, 5, "LEGENDA")
    wb.save(percorso)
    return str(percorso)
