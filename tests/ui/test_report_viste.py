"""Viste dei report (F5): lettura per tutti i ruoli, export XLSX coerente con la pagina."""

import datetime
import io
from decimal import Decimal

import openpyxl
import pytest
from django.urls import reverse

from contratti.engine.tipi import Livello
from reports.budget import LIVELLI, VISTE, Filtri, calcola_budget
from reports.piano_ordini import calcola_piano
from reports.xlsx import TIPO_XLSX

pytestmark = pytest.mark.django_db
OGGI = datetime.date.today()


@pytest.mark.parametrize("nome", ["reports:budget", "reports:piano_ordini"])
def test_anonimo_va_al_login(client, nome):
    risposta = client.get(reverse(nome))
    assert risposta.status_code == 302
    assert reverse("accounts:login") in risposta["Location"]


@pytest.mark.parametrize("nome", ["reports:budget", "reports:piano_ordini"])
def test_reader_vede_i_report(reader, dati, nome):
    assert reader.get(reverse(nome)).status_code == 200
    xlsx = reader.get(reverse(nome), {"formato": "xlsx"})
    assert xlsx.status_code == 200
    assert xlsx["Content-Type"] == TIPO_XLSX


def test_pagina_budget(reader, dati):
    html = reader.get(
        reverse("reports:budget"), {"anno": 2026, "per": "fornitore"}
    ).content.decode()
    assert "Budget 2026" in html and "IVA esclusa" in html
    assert "10.000,00 €" in html  # impegno, cassa e competenza del contratto annuale
    assert str(dati["fornitore"]) in html


def test_anno_non_valido(reader, dati):
    risposta = reader.get(reverse("reports:budget"), {"anno": "duemila"})
    assert risposta.status_code == 200
    assert "budget" not in risposta.context or risposta.context.get("budget") is None


def _righe(foglio, da=7):
    return list(foglio.iter_rows(min_row=da, values_only=True))


def test_export_budget_coincide_con_il_calcolo(reader, dati):
    risposta = reader.get(reverse("reports:budget"), {"anno": 2027, "formato": "xlsx"})
    # Senza filtro la pagina mostra l'OPEX: OPEX e CAPEX sono budget distinti.
    assert "_Budget_2027_OPEX.xlsx" in risposta["Content-Disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(risposta.content))
    assert wb.sheetnames[:3] == ["Riepilogo", "Per conto", "Dettaglio voci"]
    totali = calcola_budget(Filtri(anno=2027, natura="OPEX"), OGGI).totali()
    attese = [
        (
            vista,
            livello.value,
            *[totali[p][vista][livello] for p in ("consolidato", "previsionale", "totale")],
        )
        for vista in VISTE
        for livello in LIVELLI
    ]
    lette = _righe(wb["Riepilogo"])
    assert len(lette) == len(attese)
    for letta, attesa in zip(lette, attese, strict=True):
        assert letta[0] == "OPEX"
        assert [Decimal(str(x)) for x in letta[3:]] == list(attesa[2:])
    # Rinnovo 2027 (10.000) + previsto manuale (5.000).
    assert totali["previsionale"]["impegno"][Livello.FISSO] == Decimal("15000.00")


def test_export_piano_ordini(reader, dati):
    risposta = reader.get(reverse("reports:piano_ordini"), {"anno": 2027, "formato": "xlsx"})
    wb = openpyxl.load_workbook(io.BytesIO(risposta.content))
    piano = calcola_piano(2027, OGGI)
    sezioni = piano.sezioni()
    assert [s.natura for s in sezioni] == [n for n in ("OPEX", "CAPEX") if n in wb.sheetnames]
    for sezione in sezioni:  # un foglio per natura: righe della natura più il totale
        righe = _righe(wb[sezione.natura])
        assert len(righe) == sezione.numero_righe + 1
        assert righe[-1][0] == f"Totale {sezione.natura}"
        t = sezione.totali
        assert [Decimal(str(x)) for x in righe[-1][15:21]] == [
            *[t.contrattuale[liv] for liv in LIVELLI],
            *[t.normalizzato[liv] for liv in LIVELLI],
        ]
    assert isinstance(_righe(wb["OPEX"])[0][7], datetime.datetime)  # date come date, non testo


def test_export_piano_ordini_semplificato(reader, dati):
    risposta = reader.get(
        reverse("reports:piano_ordini"), {"anno": 2027, "formato": "xlsx", "vista": "semplificata"}
    )
    assert "Piano_ordini_2027_semplificato.xlsx" in risposta["Content-Disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(risposta.content))
    piano = calcola_piano(2027, OGGI)
    sezioni = piano.sezioni()
    assert sezioni
    for sezione in sezioni:
        foglio = wb[sezione.natura]
        intestazione = [c.value for c in foglio[6]]
        for tolta in (
            "Procedura presunta", "Lead time (mesi)", "Data limite di avvio",
            "Fase", "In ritardo", "Contrattuale fisso", "Contrattuale probabile",
            "Normalizzato 2027 fisso", "Normalizzato 2027 probabile",
        ):  # fmt: skip
            assert tolta not in intestazione
        assert (
            "Contrattuale massimo" in intestazione and "Normalizzato 2027 massimo" in intestazione
        )
        righe = _righe(foglio)
        assert len(righe) == sezione.numero_righe + 1
        assert all(len(r) == len(intestazione) for r in righe)
        t = sezione.totali
        i = intestazione.index("Contrattuale massimo")
        assert [Decimal(str(x)) for x in righe[-1][i : i + 2]] == [
            t.contrattuale[Livello.MASSIMO],
            t.normalizzato[Livello.MASSIMO],
        ]


def test_vista_piano_ordini_semplificata(reader, dati):
    url = reverse("reports:piano_ordini")
    completa = reader.get(url, {"anno": 2027}).content.decode()
    html = reader.get(url, {"anno": 2027, "vista": "semplificata"}).content.decode()
    for colonna in ("Procedura</th>", "Data limite</th>"):
        assert colonna in completa and colonna not in html
    assert "lead time" not in html and "in ritardo" not in html.split("<table")[1]
    assert "Vista completa" in html and "Vista semplificata" in completa
    assert 'name="vista" value="semplificata"' in html  # sopravvive a «Aggiorna»


def test_budget_opex_di_default_e_nature_separate(reader, dati):
    url = reverse("reports:budget")
    html = reader.get(url, {"anno": 2026}).content.decode()
    assert "Budget OPEX 2026" in html and "Budget CAPEX" not in html
    html = reader.get(url, {"anno": 2026, "natura": "separate"}).content.decode()
    assert "Budget OPEX 2026" in html and "Budget CAPEX 2026" in html
    assert "i totali non si sommano" in html
    # Il fatturato è la prima vista.
    assert html.index("Fatturato (data fattura)") < html.index("Ordinato (contratti emessi)")
