"""Import delle tabelle SAP (F11.1, SPEC §7 Rev13): richiedenti, gruppi merci, gruppi acquisti.

Codici di fantasia.
"""

import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import CommandError, call_command
from django.urls import reverse

from accounts.roles import ADMIN, EDITOR
from contratti import models as m
from contratti.importer.tabelle import FileNonValido
from contratti.importer.tabelle_sap import importa_tabelle_sap

pytestmark = pytest.mark.django_db


def _xlsx(tmp_path, fogli: dict[str, list[list]], nome="tabelle_sap.xlsx"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for titolo, righe in fogli.items():
        ws = wb.create_sheet(titolo)
        for r in righe:
            ws.append(r)
    percorso = tmp_path / nome
    wb.save(percorso)
    return percorso


COMPLETO = {
    "Richiedenti": [
        ["Codice", "Descrizione", "Predefinito"],
        [999901, "Ufficio di prova", "sì"],
        ["999902", "Altro ufficio", ""],
    ],
    "Gruppi merci": [
        ["codice", "descrizione", "predefinito", "categoria"],
        ["GM-TEST01", "Software", "x", "Licenze software"],
        ["GM-TEST02", "Servizi", None, None],
    ],
    "Gruppi acquisti": [["codice", "nome", "predefinito"], ["G01", "Gruppo di prova", "si"]],
}


def test_carica_aggiorna_idempotente(tmp_path):
    report = importa_tabelle_sap(_xlsx(tmp_path, COMPLETO))
    assert report.conteggi["richiedenti_creati"] == 2
    assert report.conteggi["gruppi_merci_creati"] == 2
    assert report.conteggi["gruppi_acquisti_creati"] == 1
    assert m.RichiedenteSAP.predefinito_attivo().codice == "999901"
    assert m.GruppoAcquisti.objects.get(codice="G01").descrizione == "Gruppo di prova"
    gm = m.GruppoMerci.objects.get(codice="GM-TEST01")
    assert gm.predefinito

    report = importa_tabelle_sap(_xlsx(tmp_path, COMPLETO))
    assert report.conteggi["richiedenti_invariati"] == 2
    assert report.conteggi["gruppi_merci_invariati"] == 2
    assert m.RichiedenteSAP.objects.count() == 2


def test_nuovo_predefinito_e_nessuna_cancellazione(tmp_path):
    importa_tabelle_sap(_xlsx(tmp_path, COMPLETO))
    solo = {"Richiedenti": [["codice", "descrizione", "predefinito"], ["999902", "Altro", "sì"]]}
    report = importa_tabelle_sap(_xlsx(tmp_path, solo))
    assert report.conteggi["richiedenti_aggiornati"] == 1
    assert m.RichiedenteSAP.predefinito_attivo().codice == "999902"
    assert m.RichiedenteSAP.objects.filter(predefinito=True).count() == 1
    assert m.RichiedenteSAP.objects.filter(codice="999901").exists()
    # Un foglio senza predefiniti lascia quello registrato.
    importa_tabelle_sap(_xlsx(tmp_path, {"Richiedenti": [["codice"], ["999903"]]}))
    assert m.RichiedenteSAP.predefinito_attivo().codice == "999902"
    assert m.GruppoAcquisti.predefinito_attivo().codice == "G01"


def test_dry_run_non_scrive(tmp_path):
    report = importa_tabelle_sap(_xlsx(tmp_path, COMPLETO), dry_run=True)
    assert report.conteggi["richiedenti_creati"] == 2
    assert not m.RichiedenteSAP.objects.exists()


@pytest.mark.parametrize(
    ("fogli", "messaggio"),
    [
        (
            {"Richiedenti": [["codice", "predefinito"], ["1", "sì"], ["2", "sì"]]},
            "un solo predefinito",
        ),
        ({"Richiedenti": [["codice"], ["1"], [1]]}, "codice ripetuto"),
        ({"Richiedenti": [["codice", "predefinito"], ["1", "forse"]]}, "non leggibile"),
        ({"Richiedenti": [["codice", "descrizione"], [None, "Senza codice"]]}, "codice mancante"),
        ({"Richiedenti": [["descrizione"], ["Senza colonna"]]}, "colonna mancante"),
    ],
)
def test_errori_bloccanti(tmp_path, fogli, messaggio):
    corretto = {"Gruppi acquisti": [["codice"], ["G09"]]}
    with pytest.raises(FileNonValido, match=messaggio):
        importa_tabelle_sap(_xlsx(tmp_path, corretto | fogli))
    assert not m.GruppoAcquisti.objects.exists()  # tutto o niente


def test_nessun_foglio_riconosciuto(tmp_path):
    with pytest.raises(FileNonValido, match="Nessun foglio riconosciuto"):
        importa_tabelle_sap(_xlsx(tmp_path, {"Altro": [["codice"], ["1"]]}))


def test_comando(tmp_path, capsys):
    call_command("import_tabelle_sap", str(_xlsx(tmp_path, COMPLETO)), "--dry-run")
    assert "[dry-run] Richiedenti: creati 2" in capsys.readouterr().out
    call_command("import_tabelle_sap", str(_xlsx(tmp_path, COMPLETO)))
    assert "Gruppi acquisti: creati 1" in capsys.readouterr().out
    with pytest.raises(CommandError, match="un solo predefinito"):
        fogli = {"Richiedenti": [["codice", "predefinito"], ["1", "sì"], ["2", "sì"]]}
        call_command("import_tabelle_sap", str(_xlsx(tmp_path, fogli)))


def test_dall_interfaccia_solo_admin(client, make_user):
    make_user("u_editor", role=EDITOR)
    make_user("u_admin", role=ADMIN)
    client.login(username="u_editor", password="password-locale-di-prova")
    url = reverse("contratti:importazione_modello", args=["tabelle_sap"])
    assert client.get(url).status_code == 404  # tipo non disponibile per l'Editor
    client.login(username="u_admin", password="password-locale-di-prova")
    wb = openpyxl.load_workbook(io.BytesIO(client.get(url).content))
    assert wb.sheetnames == [
        "Richiedenti", "Gruppi merci", "Gruppi acquisti", "Centri di costo", "Istruzioni",
    ]  # fmt: skip
    wb["Richiedenti"].append(["999901", "Ufficio di prova", "sì"])
    buffer = io.BytesIO()
    wb.save(buffer)
    dati = {"tabelle_sap-file": SimpleUploadedFile("tabelle.xlsx", buffer.getvalue())}
    client.post(reverse("contratti:importazione_carica", args=["tabelle_sap"]), dati)
    imp = m.Importazione.objects.latest("pk")
    assert imp.anteprima["conteggi"]["richiedenti_creati"] == 1
    assert not m.RichiedenteSAP.objects.exists()
    client.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    assert m.RichiedenteSAP.predefinito_attivo().codice == "999901"


# --- Elenchi di SAP così come arrivano (gruppi merci, codici di rilascio) -------------------

GRUPPI_MERCI_SAP = {
    "Foglio1": [
        ["Gruppo Merci", "Definizione", "Definizione"],
        ["ZZ01AA01", "SOFTWARE PROVA", "SOFTWARE APPLICATIVO DI PROVA"],
        ["ZZ01AA02", "SERVIZI PROVA", None],
    ]
}
RILASCIO_SAP = {
    "Reorg prova": [
        ["Unità", "Richiedente", "Gr Acquisti Servizi", "Responsabile", None],
        [None, None, None, None, "Nota a margine"],
        ["Direzione di prova (DIRP)", "900001", "Z01", "Nome Cognome Uno", None],
        [None, None, None, None, None],
        ["ICT, Digitalizzazione e Innovazione (ICTDINN)", "090002", "Z05", "Nome Due", None],
        ["Ufficio con lettera (UFL)", "90003A", "Z05", "Nome Tre", None],
        ["Ufficio senza gruppo", "900004", None, "Nome Quattro", None],
    ]
}


def test_gruppi_merci_dall_elenco_sap(tmp_path):
    report = importa_tabelle_sap(_xlsx(tmp_path, GRUPPI_MERCI_SAP))
    assert report.conteggi["gruppi_merci_creati"] == 2
    descrizioni = dict(m.GruppoMerci.objects.values_list("codice", "descrizione"))
    # La definizione lunga; la breve se la lunga manca.
    assert descrizioni == {"ZZ01AA01": "SOFTWARE APPLICATIVO DI PROVA",
                           "ZZ01AA02": "SERVIZI PROVA"}  # fmt: skip
    assert not m.GruppoMerci.objects.filter(predefinito=True).exists()
    assert (
        importa_tabelle_sap(_xlsx(tmp_path, GRUPPI_MERCI_SAP)).conteggi["gruppi_merci_invariati"]
        == 2
    )


def test_codici_di_rilascio(tmp_path):
    report = importa_tabelle_sap(_xlsx(tmp_path, RILASCIO_SAP))
    c = report.conteggi
    assert (c["richiedenti_creati"], c["gruppi_acquisti_creati"]) == (4, 2)
    assert c["richiedenti_collegati"] == 3
    r = {x.codice: x for x in m.RichiedenteSAP.objects.select_related("gruppo_acquisti")}
    # Codici come testo: zeri iniziali e lettere restano.
    assert set(r) == {"900001", "090002", "90003A", "900004"}
    assert r["090002"].descrizione == "ICT, Digitalizzazione e Innovazione (ICTDINN)"
    assert r["090002"].gruppo_acquisti.codice == "Z05" and r["900004"].gruppo_acquisti is None
    # Un gruppo condiviso da due unità prende il nome della prima.
    assert m.GruppoAcquisti.objects.get(codice="Z05").descrizione.endswith("(ICTDINN)")
    # La riga ICTDINN è il predefinito; il responsabile non si importa.
    assert m.RichiedenteSAP.predefinito_attivo().codice == "090002"
    assert m.GruppoAcquisti.predefinito_attivo().codice == "Z05"
    assert "Nome" not in str(list(m.RichiedenteSAP.objects.values()))
    report = importa_tabelle_sap(_xlsx(tmp_path, RILASCIO_SAP))
    assert report.conteggi["richiedenti_invariati"] == 4
    assert report.conteggi["richiedenti_collegati"] == 0


def test_rilascio_non_cambia_un_predefinito_gia_scelto(tmp_path):
    m.RichiedenteSAP.objects.create(codice="999901", predefinito=True)
    importa_tabelle_sap(_xlsx(tmp_path, RILASCIO_SAP))
    assert m.RichiedenteSAP.predefinito_attivo().codice == "999901"
    assert m.GruppoAcquisti.predefinito_attivo().codice == "Z05"  # quella tabella era vuota


def test_rilascio_codice_numerico_e_richiedente_mancante(tmp_path):
    righe = [list(r) for r in RILASCIO_SAP["Reorg prova"]]
    righe[2][1] = 900001  # cella numerica: gli zeri iniziali potrebbero essersi persi
    report = importa_tabelle_sap(_xlsx(tmp_path, {"Reorg prova": righe}))
    assert [(a.codice, a.riga) for a in report.anomalie] == [("CODICE_NUMERICO", 3)]
    righe[4][1] = None
    with pytest.raises(FileNonValido, match="riga 5: richiedente mancante"):
        importa_tabelle_sap(_xlsx(tmp_path, {"Reorg prova": righe}, "mancante.xlsx"))


def test_comando_con_gli_elenchi_sap(tmp_path, capsys):
    call_command("import_tabelle_sap", str(_xlsx(tmp_path, RILASCIO_SAP | GRUPPI_MERCI_SAP)))
    uscita = capsys.readouterr().out
    assert "Gruppi merci: creati 2" in uscita and "Richiedenti: creati 4" in uscita
    assert "legati al gruppo acquisti della loro unità: 3" in uscita


def test_richiesta_senza_gruppo_acquisti_prende_quello_del_richiedente():
    from contratti import forms_rda

    g2 = m.GruppoAcquisti.objects.create(codice="G02")
    r = m.RichiedenteSAP.objects.create(codice="999902", gruppo_acquisti=g2)
    form = forms_rda.RichiestaForm(
        {"descrizione": "Richiesta di prova", "richiedente": r.pk, "gruppo_acquisti": ""},
        instance=m.RichiestaAcquisto(),
    )
    assert form.is_valid(), form.errors
    assert form.cleaned_data["gruppo_acquisti"] == g2
    assert f'"{r.pk}": {g2.pk}' in form.fields["richiedente"].widget.attrs["data-gruppi-acquisti"]
