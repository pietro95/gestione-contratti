"""Piano ordini (SPEC §5.4): sezioni OPEX/CAPEX, sottocluster, fonte, valore normalizzato.

Contratti già registrati con competenza nell'anno entrano accanto ai previsti, con il valore
contrattuale (intera durata) e il normalizzato sull'anno (competenza a giorni, come TC-11).
"""

import datetime
import io
from decimal import Decimal

import openpyxl
import pytest
from django.http import HttpResponse
from django.test import Client
from django.urls import reverse

from accounts.roles import ADMIN, READER
from contratti import models as m
from contratti.engine.tipi import Livello
from contratti.forms import FabbisognoForm
from reports.export import piano_xlsx
from reports.piano_ordini import FONTE_ORDINARIA, SENZA_SOTTOCLUSTER, calcola_piano

from .. import factories as f

pytestmark = pytest.mark.django_db
D = datetime.date
OGGI = D(2026, 6, 1)
FISSO = Livello.FISSO
PASSWORD = "password-locale-di-prova"


def _registrato(descrizione="Servizio registrato", *, inizio=None, **kw):
    """12 mesi, 1.200,00 fisso; senza previsti (decisione «dismettere»)."""
    inizio = inizio or D(2026, 7, 1)
    fab = f.fabbisogno(descrizione, decisione="dismettere", **kw.pop("fab", {}))
    c = f.contratto(fabbisogno=fab, data_inizio=inizio, durata_mesi=12, **kw)
    f.piano(f.voce(c, descrizione="Canone", importo_fisso=Decimal("1200.00")))
    return c


def _riga(piano, contratto):
    return next(r for r in piano.righe if r.contratto == contratto)


def test_registrato_con_valore_contrattuale_e_normalizzato():
    c = _registrato()  # 01/07/2026–30/06/2027: 184 giorni nel 2026 su 365
    r = _riga(calcola_piano(2026, OGGI), c)
    assert r.fase == "registrato" and r.data_limite is None
    assert r.importi[FISSO] == Decimal("1200.00")
    assert r.normalizzato[FISSO] == Decimal("604.93")
    assert r.per_natura["OPEX"][FISSO] == Decimal("1200.00")
    # Nell'anno dopo, il resto: i due anni quadrano con il valore contrattuale.
    r27 = _riga(calcola_piano(2027, OGGI), c)
    assert r27.normalizzato[FISSO] == Decimal("595.07")
    assert r.normalizzato[FISSO] + r27.normalizzato[FISSO] == r.importi[FISSO]


def test_registrato_fuori_dai_totali_contrattuali_ma_nelle_sezioni():
    c = _registrato()
    piano = calcola_piano(2026, OGGI)
    assert _riga(piano, c) not in piano.da_avviare
    assert piano.totali()[FISSO] == Decimal("0.00")
    (opex,) = piano.sezioni()
    assert opex.totali.contrattuale[FISSO] == Decimal("0.00")
    assert opex.totali.normalizzato[FISSO] == Decimal("604.93")


def test_stato_in_stipula_e_stipulato():
    futuro = _registrato("Futuro")  # avvio 01/07/2026, nessuna stipula: In stipula
    firmato = _registrato("Firmato", data_stipula=D(2026, 5, 15))
    avviato = _registrato("Avviato", inizio=D(2026, 3, 1))  # avvio già passato: Stipulato
    stipula_futura = _registrato("Stipula futura", data_stipula=D(2026, 6, 30))
    piano = calcola_piano(2026, OGGI)
    assert _riga(piano, futuro).stato == "In stipula"
    assert _riga(piano, firmato).stato == "Stipulato"
    assert _riga(piano, avviato).stato == "Stipulato"
    assert _riga(piano, stipula_futura).stato == "In stipula"


def test_esclusi_senza_competenza_senza_contratto_e_archiviati():
    scaduto = _registrato("Scaduto", inizio=D(2024, 1, 1))  # nessuna competenza nel 2026
    senza = _registrato("Senza contratto", senza_contratto=True)
    archiviato = _registrato("Archiviato", archiviato=True)
    senza_data = _registrato("Senza data")
    m.Contratto.objects.filter(pk=senza_data.pk).update(data_inizio=None)
    in_piano = _registrato("In piano")
    piano = calcola_piano(2026, OGGI)
    assert [r.contratto for r in piano.righe] == [in_piano]
    assert {scaduto, senza, archiviato, senza_data}.isdisjoint(r.contratto for r in piano.righe)
    # Lo stesso contratto scaduto pesa nel 2024: l'esclusione dipende dall'anno.
    assert _riga(calcola_piano(2024, OGGI), scaduto).normalizzato[FISSO] == Decimal("1200.00")


def test_una_tantum_registrata_entra_se_pesa():
    c = _registrato("Acquisto una tantum", una_tantum=True)
    assert _riga(calcola_piano(2026, OGGI), c).contratto.una_tantum


def test_previsto_normalizzato_sull_anno():
    fab = f.fabbisogno("Da rinnovare", decisione="rinnovare")
    c = f.contratto(
        fabbisogno=fab, data_inizio=D(2026, 1, 1), durata_mesi=12, lead_time_mesi_override=2
    )
    f.piano(f.voce(c, importo_fisso=Decimal("1200.00")))
    # Il previsto parte il 01/01/2027, data limite 31/10/2026: nel 2026 da avviare, ma senza
    # competenza nell'anno; nel 2027 da concludere, con tutta la competenza.
    p26 = next(r for r in calcola_piano(2026, OGGI).righe if r.numero_previsto == 1)
    assert p26.da_avviare and p26.stato == "Previsto"
    assert p26.importi[FISSO] > 0 and p26.normalizzato[FISSO] == Decimal("0.00")
    p27 = next(r for r in calcola_piano(2027, OGGI).righe if r.numero_previsto == 1)
    assert p27.fase == "da_concludere"
    assert p27.normalizzato[FISSO] == p27.importi[FISSO]
    # Il registrato dello stesso fabbisogno è un'altra riga, nel piano 2026.
    assert _riga(calcola_piano(2026, OGGI), c).normalizzato[FISSO] == Decimal("1200.00")


@pytest.fixture
def struttura():
    it, tlc = f.cluster("IT"), f.cluster("TLC")
    reti = m.Sottocluster.objects.create(cluster=it, nome="Reti")
    server = m.Sottocluster.objects.create(cluster=it, nome="Server")
    wbs = m.WBS.objects.create(
        codice_wbs="X.0001", descrizione="Investimenti IT", tipo="cluster", cluster=it
    )
    return {"it": it, "tlc": tlc, "reti": reti, "server": server, "wbs": wbs}


def _misto(struttura, descrizione="Piattaforma mista", **kw):
    fab = f.fabbisogno(
        descrizione, cluster=struttura["it"], sottocluster=struttura["reti"],
        decisione="dismettere", **kw,
    )  # fmt: skip
    c = f.contratto(fabbisogno=fab, data_inizio=D(2026, 1, 1), durata_mesi=12)
    f.piano(f.voce(c, descrizione="Supporto", importo_fisso=Decimal("600.00")))
    f.piano(
        f.voce(
            c, descrizione="Licenze", natura="CAPEX", wbs=struttura["wbs"],
            importo_fisso=Decimal("1000.00"),
        ),
        modalita=m.ModalitaFatturazione.UNA_TANTUM, periodicita="",
    )  # fmt: skip
    return c


def test_sezioni_opex_e_capex_con_quota_per_natura(struttura):
    c = _misto(struttura)
    piano = calcola_piano(2026, OGGI)
    opex, capex = piano.sezioni()
    assert (opex.natura, capex.natura) == ("OPEX", "CAPEX")
    # Un contratto misto compare in entrambe, con la sola quota della natura.
    (ro,) = opex.cluster[0].sottocluster[0].righe
    (rc,) = capex.cluster[0].sottocluster[0].righe
    assert ro.riga.contratto == rc.riga.contratto == c
    assert ro.importi[FISSO] == Decimal("600.00") and rc.importi[FISSO] == Decimal("1000.00")
    assert ro.normalizzato[FISSO] == Decimal("600.00")
    assert rc.normalizzato[FISSO] == Decimal("1000.00")
    # Totali propri, mai sommati.
    assert opex.totali.normalizzato[FISSO] == Decimal("600.00")
    assert capex.totali.normalizzato[FISSO] == Decimal("1000.00")
    assert [s.natura for s in piano.sezioni("CAPEX")] == ["CAPEX"]
    assert piano.sezioni("CAPEX")[0].numero_righe == 1


def test_raggruppamento_cluster_sottocluster_con_subtotali(struttura):
    _misto(struttura, "Rete A")
    _misto(struttura, "Rete B")
    _registrato("Server C", fab={"cluster": struttura["it"], "sottocluster": struttura["server"]})
    _registrato("Senza sottocluster", fab={"cluster": struttura["it"]})
    _registrato("Telefonia", fab={"cluster": struttura["tlc"]})
    opex = calcola_piano(2026, OGGI).sezioni("OPEX")[0]
    assert [g.cluster.codice for g in opex.cluster] == ["IT", "TLC"]
    it = opex.cluster[0]
    assert [g.nome for g in it.sottocluster] == ["Reti", "Server", SENZA_SOTTOCLUSTER]
    reti = it.sottocluster[0]
    assert len(reti.righe) == 2 and reti.totali.normalizzato[FISSO] == Decimal("1200.00")
    # Il totale del cluster è la somma dei subtotali, quello della sezione dei cluster.
    assert it.totali.normalizzato[FISSO] == sum(
        (g.totali.normalizzato[FISSO] for g in it.sottocluster), Decimal("0")
    )
    assert opex.totali.normalizzato[FISSO] == sum(
        (g.totali.normalizzato[FISSO] for g in opex.cluster), Decimal("0")
    )
    assert opex.numero_righe == 5


def test_fonte_da_fabbisogno_poi_wbs_poi_ordinaria(struttura):
    pnrr = m.FonteFinanziamento.objects.create(nome="Fonte A (prova)")
    altra = m.FonteFinanziamento.objects.create(nome="Fonte B (prova)")
    struttura["wbs"].fonte_finanziamento = altra
    struttura["wbs"].save()
    c = _misto(struttura, "Con fonte propria", fonte_finanziamento=pnrr)
    senza = _misto(struttura, "Solo WBS")
    ordinario = _registrato("Ordinario")
    piano = calcola_piano(2026, OGGI)
    # La fonte del fabbisogno vale per tutte le sue quote, anche per la WBS con fonte.
    assert _riga(piano, c).fonti == {"OPEX": "Fonte A (prova)", "CAPEX": "Fonte A (prova)"}
    # Senza fonte sul fabbisogno: quella della WBS, solo per la quota su WBS.
    assert _riga(piano, senza).fonti == {"OPEX": FONTE_ORDINARIA, "CAPEX": "Fonte B (prova)"}
    assert _riga(piano, ordinario).fonte == FONTE_ORDINARIA


def test_fonte_del_previsto_manuale_dalla_wbs(struttura):
    fonte = m.FonteFinanziamento.objects.create(nome="Fonte C (prova)")
    struttura["wbs"].fonte_finanziamento = fonte
    struttura["wbs"].save()
    f.fabbisogno(
        "Nuovo investimento", cluster=struttura["it"], natura_previsti="CAPEX",
        wbs_previsti=struttura["wbs"], previsto_data_avvio=D(2027, 3, 1), previsto_durata_mesi=12,
    )  # fmt: skip
    r = next(r for r in calcola_piano(2026, OGGI).righe if r.stato == "Previsto")
    assert r.fonti == {"CAPEX": "Fonte C (prova)"}


def test_sottocluster_validato_sul_cluster(struttura):
    dati = {
        "descrizione": "Fabbisogno di prova",
        "cluster": struttura["tlc"].pk,
        "categoria": f.categoria().pk,
        "decisione": "da_valutare",
        "sottocluster": struttura["reti"].pk,  # sottocluster di IT su un fabbisogno TLC
    }
    form = FabbisognoForm(dati)
    assert not form.is_valid() and "sottocluster" in form.errors
    dati["cluster"] = struttura["it"].pk
    assert FabbisognoForm(dati).is_valid()


@pytest.fixture
def lettore(make_user):
    make_user("lettore_piano", role=READER)
    client = Client()
    client.login(username="lettore_piano", password=PASSWORD)
    return client


def test_pagina_con_sezioni_sottocluster_e_fonte(struttura, lettore):
    fonte = m.FonteFinanziamento.objects.create(nome="Fonte D (prova)")
    _misto(struttura, fonte_finanziamento=fonte)
    url = reverse("reports:piano_ordini")
    risposta = lettore.get(url, {"anno": 2026, "area": "tutti"})
    html = risposta.content.decode()
    assert 'id="sezione-opex"' in html and 'id="sezione-capex"' in html
    assert "Reti" in html and "Fonte D (prova)" in html and "Normalizzato 2026" in html
    assert [s.natura for s in risposta.context["sezioni"]] == ["OPEX", "CAPEX"]
    solo = lettore.get(url, {"anno": 2026, "area": "tutti", "natura": "CAPEX"})
    assert [s.natura for s in solo.context["sezioni"]] == ["CAPEX"]
    assert 'id="sezione-opex"' not in solo.content.decode()


def test_xlsx_un_foglio_per_natura(struttura, lettore):
    fonte = m.FonteFinanziamento.objects.create(nome="Fonte E (prova)")
    _misto(struttura, fonte_finanziamento=fonte)
    risposta = lettore.get(
        reverse("reports:piano_ordini"), {"anno": 2026, "area": "tutti", "formato": "xlsx"}
    )
    wb = openpyxl.load_workbook(io.BytesIO(risposta.content))
    assert wb.sheetnames[:2] == ["OPEX", "CAPEX"]
    intestazione = [c.value for c in wb["OPEX"][6]]
    righe = {n: list(wb[n].iter_rows(min_row=7, values_only=True)) for n in ("OPEX", "CAPEX")}
    col = {nome: i for i, nome in enumerate(intestazione)}
    for natura, valore in (("OPEX", "600.00"), ("CAPEX", "1000.00")):
        riepilogo, riga, totale = righe[natura]  # il contratto stipulato è raggruppato
        assert riepilogo[0] == "1 contratti già stipulati"
        assert riepilogo[col["Normalizzato 2026 fisso"]] == riga[col["Normalizzato 2026 fisso"]]
        assert riga[col["Sottocluster"]] == "Reti" and riga[col["Cluster"]] == "IT"
        assert riga[col["Fonte di finanziamento"]] == "Fonte E (prova)"
        assert riga[col["Stato"]] == "Stipulato"
        assert Decimal(str(riga[col["Contrattuale fisso"]])) == Decimal(valore)
        assert Decimal(str(riga[col["Normalizzato 2026 fisso"]])) == Decimal(valore)
        assert totale[0] == f"Totale {natura}"
        assert Decimal(str(totale[col["Normalizzato 2026 fisso"]])) == Decimal(valore)


def test_stipulati_in_una_riga_espandibile_per_sottocluster(make_user):
    a = _registrato("Stipulato A", inizio=D(2026, 3, 1))
    b = _registrato("Stipulato B", inizio=D(2026, 2, 1))
    _registrato("In stipula")  # avvio futuro, nessuna stipula: resta in elenco
    piano = calcola_piano(2026, OGGI)
    (opex,) = piano.sezioni()
    (gc,) = opex.cluster
    (gs,) = gc.sottocluster
    assert {rn.riga.contratto for rn in gs.stipulati} == {a, b}
    assert [rn.riga.contratto.fabbisogno.descrizione for rn in gs.altre] == ["In stipula"]
    somma = gs.riepilogo_stipulati
    assert somma.contrattuale[FISSO] == Decimal("2400.00")
    assert somma.normalizzato[FISSO] == sum(
        (_riga(piano, c).normalizzato[FISSO] for c in (a, b)), Decimal("0.00")
    )
    # Il raggruppamento è solo di vista: subtotali e totali contano tutte le righe.
    assert opex.totali.normalizzato[FISSO] == gs.totali.normalizzato[FISSO]
    assert gs.totali.normalizzato[FISSO] > somma.normalizzato[FISSO]

    make_user("admin_piano", role=ADMIN)  # solo l'Admin può simulare la data di oggi
    admin = Client()
    admin.login(username="admin_piano", password=PASSWORD)
    html = admin.get(
        reverse("reports:piano_ordini"), {"anno": 2026, "area": "tutti", "oggi": "2026-06-01"}
    ).content.decode()
    assert html.count("2 contratti già stipulati") == 1
    assert html.count('class="collapse stip-opex-') == 2  # i due dettagli, nascosti di default
    assert html.count('data-bs-target=".stip-opex-') == 1


def test_xlsx_stipulati_raggruppati_con_le_righe_comprimibili(make_user):
    _registrato("Stipulato A", inizio=D(2026, 3, 1))
    _registrato("Stipulato B", inizio=D(2026, 2, 1))
    piano = calcola_piano(2026, OGGI)
    wb = piano_xlsx(piano, OGGI)
    ws = wb["OPEX"]
    assert ws.sheet_properties.outlinePr.summaryBelow is False  # pulsante sul riepilogo
    righe = list(ws.iter_rows(min_row=7, values_only=True))
    assert [r[0] for r in righe][0] == "2 contratti già stipulati"
    livelli = [ws.row_dimensions[i].outlineLevel for i in range(7, 7 + len(righe))]
    nascoste = [ws.row_dimensions[i].hidden for i in range(7, 7 + len(righe))]
    assert livelli == [0, 1, 1, 0]  # riepilogo, due dettagli, totale
    assert nascoste == [False, True, True, False]
    # Il riepilogo somma i dettagli; il totale della sezione non li conta due volte.
    assert "non sommare entrambe" in ws["A3"].value
    totale = piano.sezioni()[0].totali
    assert Decimal(str(righe[-1][-4])) == totale.normalizzato[FISSO]
    assert Decimal(str(righe[0][-4])) == Decimal(str(righe[1][-4])) + Decimal(str(righe[2][-4]))


def test_pdf_senza_totali_in_testa(make_user, monkeypatch):
    from django.template.loader import render_to_string

    from reports import pdf as modulo_pdf

    make_user("admin_pdf2", role=ADMIN)
    client = Client()
    client.login(username="admin_pdf2", password=PASSWORD)
    _registrato("Stipulato A", inizio=D(2026, 3, 1))
    catturato = {}

    def finto(template, contesto, nome, request):
        catturato.update(template=template, contesto=contesto)
        return HttpResponse(b"%PDF", content_type="application/pdf")

    monkeypatch.setattr(modulo_pdf, "risposta", finto)
    client.get(reverse("reports:piano_ordini") + "?anno=2026&area=tutti&formato=pdf")
    html = render_to_string(catturato["template"], catturato["contesto"], None)
    assert "Ordini da emettere" in html
    assert "valore probabile" not in html and "valore massimo" not in html
    assert "Totale OPEX" in html  # i totali restano in fondo a ogni sezione


def test_pdf_due_versioni_dei_stipulati(make_user, monkeypatch):
    from django.template.loader import render_to_string

    from reports import pdf as modulo_pdf

    make_user("admin_pdf", role=ADMIN)
    client = Client()
    client.login(username="admin_pdf", password=PASSWORD)
    _registrato("Stipulato A", inizio=D(2026, 3, 1))
    _registrato("Stipulato B", inizio=D(2026, 2, 1))
    catturato = {}

    def finto(template, contesto, nome, request):
        catturato.update(template=template, contesto=contesto, nome=nome)
        return HttpResponse(b"%PDF", content_type="application/pdf")

    monkeypatch.setattr(modulo_pdf, "risposta", finto)
    url = reverse("reports:piano_ordini") + "?anno=2026&area=tutti&oggi=2026-06-01&formato=pdf"
    client.get(url)
    completo = render_to_string(catturato["template"], catturato["contesto"], None)
    assert "Stipulato A" in completo and "Stipulato B" in completo
    assert catturato["nome"].endswith("_Piano_ordini_2026.pdf")
    client.get(url + "&raggruppa=1")
    raggruppato = render_to_string(catturato["template"], catturato["contesto"], None)
    assert "2 contratti già stipulati" in raggruppato
    assert "Stipulato A" not in raggruppato and "Stipulato B" not in raggruppato
    assert catturato["nome"].endswith("_Piano_ordini_2026_stipulati_raggruppati.pdf")
    html = client.get(url.replace("&formato=pdf", "")).content.decode()
    assert "PDF con i contratti stipulati raggruppati" in html


def test_pagina_senza_totali_in_testa(struttura, lettore):
    """In testa niente totali per natura: i totali stanno solo in fondo a ogni sezione."""
    _misto(struttura)
    html = lettore.get(
        reverse("reports:piano_ordini"), {"anno": 2026, "area": "tutti"}
    ).content.decode()
    assert "Ordini da emettere" in html and "Possibile frazionamento" in html
    assert "valore probabile" not in html and "valore massimo" not in html
    assert html.count("<tfoot") == 2  # il totale di OPEX e quello di CAPEX
