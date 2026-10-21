"""F9.2: CAPEX per WBS con il dettaglio per conto; pagina «Spese per WBS»."""

import datetime
import io
from decimal import Decimal

import openpyxl
import pytest
from django.test import Client
from django.urls import reverse

from accounts.roles import EDITOR, READER
from contratti import models as m
from contratti.engine.tipi import Livello
from reports.budget import Filtri, calcola_budget
from reports.export import budget_xlsx
from reports.wbs import elenco_wbs, spese_wbs

from .. import factories as f

pytestmark = pytest.mark.django_db
D = datetime.date
OGGI = D(2026, 6, 1)
PASSWORD = "password-locale-di-prova"


def _conto(codice, cluster):
    return f.conto(
        codice=codice, descrizione=f"Conto {codice}", valido_dal=D(2020, 1, 1), cluster=cluster
    )


@pytest.fixture
def capex():
    it = f.cluster("IT")
    wbs = m.WBS.objects.create(
        codice_wbs="X.0001.00.000", descrizione="Investimenti IT", tipo="cluster", cluster=it
    )
    progetto = m.WBS.objects.create(codice_wbs="X.0002.00.000", descrizione="Nuovo impianto")
    c = f.contratto(
        fabbisogno=f.fabbisogno("Piattaforma", cluster=it),
        data_inizio=D(2026, 1, 1),
        durata_mesi=24,
        numero_contratto="5800000201",
    )
    conto_a, conto_b = _conto("11000001", it), _conto("11000002", it)
    v1 = f.voce(c, natura="CAPEX", wbs=wbs, conto=conto_a, importo_fisso=Decimal("2400.00"))
    v2 = f.voce(c, natura="CAPEX", wbs=wbs, conto=conto_b, importo_fisso=Decimal("1000.01"))
    f.piano(v1, modalita="anticipata", periodicita="annuale")
    f.piano(v2, modalita="anticipata", periodicita="annuale")
    altro = f.contratto(fabbisogno=f.fabbisogno("Impianto", cluster=it), data_inizio=D(2026, 1, 1))
    v3 = f.voce(altro, natura="CAPEX", wbs=progetto, importo_fisso=Decimal("500.00"))
    f.piano(v3, modalita="anticipata", periodicita="annuale")
    return {"wbs": wbs, "progetto": progetto, "conti": (conto_a, conto_b), "contratto": c}


def test_capex_per_wbs_con_dettaglio_per_conto(capex):
    b = calcola_budget(Filtri(anno=2026, natura="CAPEX"), OGGI)
    (g,) = b.gruppi("wbs")  # la WBS a progetto è fuori dal budget del cluster
    assert g.etichetta.startswith("X.0001.00.000")
    sotto = g.sotto("conto")
    assert [s.chiave for s in sotto] == ["11000001", "11000002"]
    for vista in ("cassa", "impegno", "competenza"):
        somma = sum((s.totali["totale"][vista][Livello.FISSO] for s in sotto), Decimal("0.00"))
        assert somma == g.totali["totale"][vista][Livello.FISSO]
    wb = budget_xlsx(b, {"OPEX": "conto", "CAPEX": "wbs"})
    assert "CAPEX per WBS e conto" in wb.sheetnames


def test_spese_per_wbs_anno_per_anno_quadrano(capex):
    s = spese_wbs(capex["wbs"], "cassa", Livello.FISSO)
    assert len(s.righe) == 2 and s.anni == [2026, 2027]
    # Anticipata annuale su 24 mesi: una rata per anno. 1.000,01 / 2 = 500,005 → 500,01
    # (ROUND_HALF_UP); l'ultima rata prende il residuo: 500,00.
    assert s.totali() == {2026: Decimal("1700.01"), 2027: Decimal("1700.00")}
    assert s.totale == Decimal("3400.01")  # = valore delle voci, al centesimo
    impegno = spese_wbs(capex["wbs"], "impegno", Livello.FISSO)
    assert impegno.totali() == {2026: Decimal("3400.01")}
    progetto = spese_wbs(capex["progetto"], "cassa", Livello.FISSO)
    assert progetto.totale == Decimal("500.00")


def _client(make_user, nome, ruolo):
    make_user(nome, role=ruolo)
    c = Client()
    c.login(username=nome, password=PASSWORD)
    return c


def test_pagine_wbs_ed_export(capex, make_user):
    tlc = f.cluster("TLC")
    m.WBS.objects.create(codice_wbs="X.0003.00.000", descrizione="TLC", tipo="cluster", cluster=tlc)
    m.AssegnazioneCluster.objects.create(username="u_it", cluster=f.cluster("IT"))
    client = _client(make_user, "u_it", READER)
    html = client.get(reverse("reports:wbs")).content.decode()
    # Di default: WBS dei propri cluster e tutte le WBS a progetto.
    assert "X.0001.00.000" in html and "X.0002.00.000" in html and "X.0003.00.000" not in html
    html = client.get(reverse("reports:wbs") + "?area=tutti").content.decode()
    assert "X.0003.00.000" in html
    url = reverse("reports:wbs_spese", args=[capex["wbs"].pk])
    risposta = client.get(url + "?vista=cassa&livello=fisso")
    assert risposta.status_code == 200
    html = risposta.content.decode()
    assert "5800000201" in html and "1.700,01" in html and "3.400,01" in html
    xlsx = client.get(url + "?vista=cassa&livello=fisso&formato=xlsx")
    assert "Spese_WBS_X.0001.00.000" in xlsx["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(xlsx.content))["Spese per anno"]
    valori = [c for riga in ws.iter_rows(values_only=True) for c in riga]
    assert "Totale" in valori


def test_elenco_wbs_filtri(capex):
    assert {w.codice_wbs for w in elenco_wbs([], "progetto")} == {"X.0002.00.000"}
    assert {w.n_voci for w in elenco_wbs([]) if w.codice_wbs == "X.0001.00.000"} == {2}


def test_spese_wbs_aperte_a_tutti_i_ruoli(capex, make_user):
    client = _client(make_user, "u_editor", EDITOR)
    assert client.get(reverse("reports:wbs_spese", args=[capex["wbs"].pk])).status_code == 200
