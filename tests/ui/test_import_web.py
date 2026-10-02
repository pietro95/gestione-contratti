"""F9.1: import dall'interfaccia. Anteprima senza scritture, conferma, registro, permessi."""

import datetime
import io
from pathlib import Path

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from accounts.roles import ADMIN, EDITOR
from contratti import models as m

from .. import factories as f
from .conftest import PASSWORD

pytestmark = pytest.mark.django_db
D = datetime.date
FIXTURE = Path(__file__).resolve().parents[1] / "fixtures" / "contratti_it_sample.xlsx"
S = m.Importazione.Stato


def _file(nome, contenuto: bytes):
    return SimpleUploadedFile(nome, contenuto)


def _csv_date(*righe):
    testo = "numero_contratto;data_emissione\n" + "".join(f"{n};{d}\n" for n, d in righe)
    return _file("date.csv", testo.encode("utf-8"))


def _carica(client, tipo, file, **campi):
    dati = {f"{tipo}-file": file} | {f"{tipo}-{k}": v for k, v in campi.items()}
    return client.post(reverse("contratti:importazione_carica", args=[tipo]), dati)


def _ultimo():
    return m.Importazione.objects.latest("pk")


@pytest.fixture
def altro_editor(make_user):
    make_user("u_editor2", role=EDITOR)
    c = Client()
    c.login(username="u_editor2", password=PASSWORD)
    return c


def test_reader_non_accede(reader, dati):
    assert reader.get(reverse("contratti:importazioni")).status_code == 403
    assert _carica(reader, "date_emissione", _csv_date()).status_code == 403
    assert m.Importazione.objects.count() == 1  # solo quella della fixture


def test_editor_vede_gli_import_dei_dati_ma_non_il_piano_dei_conti(editor, dati):
    html = editor.get(reverse("contratti:importazioni")).content.decode()
    assert "Excel dei contratti" in html and "Date di emissione" in html
    assert "Quadratura con l&#x27;Excel" in html and "Piano dei conti" not in html
    conti = _file("conti.csv", b"codice;descrizione;valido_dal\n600001;Canoni;01/01/2026\n")
    assert _carica(editor, "conti", conti).status_code == 403
    assert not m.ContoCoge.objects.filter(codice="600001").exists()


def test_date_anteprima_senza_scritture_poi_conferma(editor, dati):
    c = dati["contratto"]  # numero 5800000001, senza data di emissione
    risposta = _carica(editor, "date_emissione", _csv_date(("5800000001", "15/12/2025")))
    imp = _ultimo()
    assert risposta.url == reverse("contratti:importazione", args=[imp.pk])
    assert imp.stato == S.ANTEPRIMA and imp.anteprima["conteggi"]["compilati"] == 1
    c.refresh_from_db()
    assert c.data_emissione is None  # l'anteprima non scrive

    html = editor.get(risposta.url).content.decode()
    assert "Anteprima: nessun dato è stato ancora scritto" in html
    assert "Conferma l&#x27;import" in html or "Conferma l'import" in html

    editor.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    imp.refresh_from_db()
    c.refresh_from_db()
    assert imp.stato == S.CONFERMATA and imp.esito["conteggi"]["compilati"] == 1
    assert c.data_emissione == D(2025, 12, 15)
    assert imp.confermata_da.username == "u_editor"
    # Una seconda conferma non riesegue nulla.
    editor.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    assert m.Importazione.objects.get(pk=imp.pk).esito == imp.esito

    # Stesso file di nuovo: avviso, ed è innocuo.
    _carica(editor, "date_emissione", _csv_date(("5800000001", "15/12/2025")))
    nuovo = _ultimo()
    html = editor.get(reverse("contratti:importazione", args=[nuovo.pk])).content.decode()
    assert "Lo stesso file è già stato importato" in html
    assert nuovo.anteprima["conteggi"]["invariati"] == 1


def test_conferma_solo_di_chi_ha_caricato_o_admin(editor, altro_editor, dati, come):
    _carica(editor, "date_emissione", _csv_date(("5800000001", "15/12/2025")))
    imp = _ultimo()
    altro_editor.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    imp.refresh_from_db()
    assert imp.stato == S.ANTEPRIMA
    admin = come(ADMIN)
    admin.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    imp.refresh_from_db()
    assert imp.stato == S.CONFERMATA and imp.confermata_da.username == "u_admin"


def test_anteprima_scaduta_e_annullamento(editor, dati):
    _carica(editor, "date_emissione", _csv_date(("5800000001", "15/12/2025")))
    imp = _ultimo()
    m.Importazione.objects.filter(pk=imp.pk).update(
        caricata_il=timezone.now() - datetime.timedelta(hours=25)
    )
    editor.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    imp.refresh_from_db()
    assert imp.stato == S.SCADUTA
    dati["contratto"].refresh_from_db()
    assert dati["contratto"].data_emissione is None

    _carica(editor, "date_emissione", _csv_date(("5800000001", "15/12/2025")))
    imp = _ultimo()
    editor.post(reverse("contratti:importazione_annulla", args=[imp.pk]))
    imp.refresh_from_db()
    assert imp.stato == S.ANNULLATA


def test_formato_e_file_illeggibile(editor, dati):
    risposta = _carica(editor, "date_emissione", _file("date.txt", b"x"))
    assert risposta.status_code == 302
    assert m.Importazione.objects.count() == 1  # rifiutato prima di registrarlo
    # Un .xlsx che non è uno zip si rifiuta prima di registrarlo.
    _carica(editor, "date_emissione", _file("date.xlsx", b"non sono un excel"))
    assert m.Importazione.objects.count() == 1
    _carica(editor, "date_emissione", _file("date.csv", b"numero;quando\n1;2\n"))
    imp = _ultimo()
    assert imp.stato == S.ERRORE and "Colonna mancante: numero_contratto" in imp.messaggio


def test_excel_dei_contratti_anteprima_e_conferma(editor):
    with open(FIXTURE, "rb") as fh:
        _carica(editor, "contratti", _file("Contratti_IT.xlsx", fh.read()), area="IT")
    imp = _ultimo()
    creati = imp.anteprima["conteggi"]["contratti_creati"]
    assert imp.stato == S.ANTEPRIMA and creati > 0
    assert imp.opzioni == {"area": "IT"}
    assert not m.Contratto.objects.exists()
    # Il report dell'anteprima si scarica con tutte le anomalie.
    report = editor.get(reverse("contratti:importazione_report", args=[imp.pk]))
    wb = openpyxl.load_workbook(io.BytesIO(report.content))
    assert wb.sheetnames[:2] == ["Riepilogo", "Anomalie"]
    # Due righe d'intestazione, una vuota, le colonne: le anomalie partono dalla riga 5.
    assert wb["Anomalie"].max_row - 4 == len(imp.anteprima["anomalie"])

    editor.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    imp.refresh_from_db()
    assert imp.stato == S.CONFERMATA
    assert m.Contratto.objects.count() == creati == imp.esito["conteggi"]["contratti_creati"]

    # Quadratura sullo stesso file: nessuna conferma, report pronto da scaricare.
    with open(FIXTURE, "rb") as fh:
        _carica(editor, "quadratura", _file("Contratti_IT.xlsx", fh.read()), area="IT")
    q = _ultimo()
    assert q.stato == S.COMPLETATA and q.anteprima["conteggi"]["righe_confrontate"] > 0
    assert q.anteprima["riepilogo"]
    scarico = editor.get(reverse("contratti:importazione_report", args=[q.pk]))
    assert openpyxl.load_workbook(io.BytesIO(scarico.content)).sheetnames[0] == "Riepilogo"
    html = editor.get(reverse("contratti:importazione", args=[q.pk])).content.decode()
    assert "Totali per anno" in html and "Conferma l" not in html


def test_excel_senza_area_rifiutato(editor):
    with open(FIXTURE, "rb") as fh:
        _carica(editor, "contratti", _file("Contratti_IT.xlsx", fh.read()))
    assert not m.Importazione.objects.exists()


def test_admin_carica_il_piano_dei_conti_dal_modello(come):
    admin = come(ADMIN)
    modello = admin.get(reverse("contratti:importazione_modello", args=["conti"]))
    wb = openpyxl.load_workbook(io.BytesIO(modello.content))
    assert wb.sheetnames == ["Conti", "Istruzioni"]
    wb["Conti"].append(["60000001", "Canoni di prova", datetime.datetime(2026, 1, 1), None])
    buffer = io.BytesIO()
    wb.save(buffer)
    _carica(admin, "conti", _file("conti.xlsx", buffer.getvalue()))
    imp = _ultimo()
    assert imp.anteprima["conteggi"] == {"creati": 1, "aggiornati": 0, "invariati": 0}
    admin.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    assert m.ContoCoge.objects.get(codice="60000001").valido_dal == D(2026, 1, 1)


def test_modello_delle_date_si_reimporta(editor, dati):
    modello = editor.get(reverse("contratti:importazione_modello", args=["date_emissione"]))
    wb = openpyxl.load_workbook(io.BytesIO(modello.content))
    assert [c.value for c in wb["Date"][1]] == ["numero_contratto", "data_emissione"]
    wb["Date"].append([5800000001, datetime.datetime(2025, 12, 15)])
    buffer = io.BytesIO()
    wb.save(buffer)
    _carica(editor, "date_emissione", _file("date.xlsx", buffer.getvalue()))
    assert _ultimo().anteprima["conteggi"]["compilati"] == 1
    # Il modello del piano dei conti non è per l'Editor.
    assert editor.get(reverse("contratti:importazione_modello", args=["conti"])).status_code == 404


def test_registro_in_elenco(editor, dati):
    f.contratto(numero_contratto="5800000099")
    _carica(editor, "date_emissione", _csv_date(("5800000099", "01/02/2026")))
    html = editor.get(reverse("contratti:importazioni")).content.decode()
    assert "date.csv" in html and "In anteprima" in html


def test_xlsx_che_si_espande_troppo_rifiutato(editor, dati):
    """Zip bomb: pochi KB compressi, centinaia di MB decompressi. Non si apre."""
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (210 * 1024 * 1024))
    risposta = _carica(editor, "date_emissione", _file("bomba.xlsx", buffer.getvalue()))
    assert risposta.status_code == 302
    assert m.Importazione.objects.count() == 1
