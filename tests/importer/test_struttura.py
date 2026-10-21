"""F9.2: import della struttura ICTDINN dal modello approvato (valori fittizi)."""

import datetime

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse

from accounts.roles import ADMIN, EDITOR
from contratti import importazioni as imp_svc
from contratti import models as m
from contratti.importer import modello_struttura, tabelle
from contratti.importer.struttura import importa_struttura

from .. import factories as f

pytestmark = pytest.mark.django_db
D = datetime.date
PASSWORD = "password-locale-di-prova"
COLONNE = {"Cluster": 3, "Responsabili": 3, "Conti": 6, "WBS": 6}  # senza «Controllo»


def _file(tmp_path, nome="struttura.xlsx", **fogli) -> str:
    """Il modello approvato, con le righe indicate al posto di quelle proposte."""
    wb = modello_struttura.crea()
    for titolo, n in COLONNE.items():
        ws = wb[titolo]
        for riga in ws.iter_rows(min_row=2, max_col=n):
            for cella in riga:
                cella.value = None
        for i, valori in enumerate(fogli.get(titolo.lower(), []), start=2):
            for j, v in enumerate(valori, start=1):
                ws.cell(i, j, v)
    percorso = tmp_path / nome
    wb.save(percorso)
    return str(percorso)


# IT e TLC esistono già (migration 0013); XX è un cluster nuovo, fittizio.
CLUSTER = [("IT", "IT", ""), ("TLC", "TLC", ""), ("XX", "Cluster di prova", "")]
CONTI = [
    ("60000001", "Canoni software (prova)", "IT", "01/01/2020", "31/12/2025", "60000002"),
    ("60000002", "Canoni software dal 2026 (prova)", "IT", "01/01/2026", None, None),
    ("60000003", "Telefonia (prova)", "TLC", "01/01/2020", None, None),
]
WBS = [
    ("X.0001.00.000", "Investimenti IT (prova)", "Cluster", "IT", "Sì", ""),
    ("X.0002.00.000", "Nuovo impianto (prova)", "Progetto", "", "Sì", ""),
    ("X.0003.00.000", "Dismessa (prova)", "Progetto", "", "No", ""),
]
RESPONSABILI = [("m.rossi", "IT", ""), ("M.Rossi", "TLC", ""), ("a.bianchi", "TUTTI", "")]


def _completo(tmp_path, **altro):
    fogli = {"cluster": CLUSTER, "conti": CONTI, "wbs": WBS, "responsabili": RESPONSABILI}
    return _file(tmp_path, **(fogli | altro))


def test_modello_generato_come_quello_approvato():
    """Fogli, intestazioni e colonna «Controllo» del modello Rev00 (approvato da Pietro)."""
    wb = modello_struttura.crea()
    assert wb.sheetnames == ["Istruzioni", "Cluster", "Responsabili", "Conti", "WBS"]
    intestazioni = {
        "Cluster": ("codice", "nome", "note", "Controllo"),
        "Responsabili": ("utente", "cluster", "note", "Controllo"),
        "Conti": ("codice", "descrizione", "cluster", "valido_dal", "valido_al",
                  "sostituito_da", "Controllo"),
        "WBS": ("codice_wbs", "descrizione", "tipo", "cluster", "attiva", "note", "Controllo"),
    }  # fmt: skip
    for foglio, attese in intestazioni.items():
        ws = wb[foglio]
        assert tuple(c.value for c in ws[1]) == attese
        assert str(ws.cell(2, len(attese)).value).startswith("=IF(")
        assert ws.max_row == 501
    assert [wb["Cluster"].cell(i, 1).value for i in range(2, 7)] == [
        "IT", "TLC", "OT", "SICFIS", "INN"
    ]  # fmt: skip


def test_import_completo_e_idempotente(tmp_path):
    percorso = _completo(tmp_path)
    anteprima = importa_struttura(percorso, dry_run=True)
    assert anteprima.conteggi["conti_creati"] == 3
    assert not m.ContoCoge.objects.exists()  # l'anteprima non scrive
    r = importa_struttura(percorso)
    assert r.conteggi["cluster_creati"] == 1
    assert (r.conteggi["conti_creati"], r.conteggi["wbs_create"]) == (3, 3)
    assert r.conteggi["assegnazioni_create"] == 3
    vecchio = m.ContoCoge.objects.get(codice="60000001")
    assert [c.codice for c in vecchio.clusters.all()] == ["IT"] and vecchio.valido_al == D(
        2025, 12, 31
    )
    assert vecchio.sostituito_da.codice == "60000002"
    assert m.WBS.objects.get(codice_wbs="X.0001.00.000").tipo == m.WBS.Tipo.CLUSTER
    assert m.WBS.objects.get(codice_wbs="X.0003.00.000").archiviato
    assert set(m.AssegnazioneCluster.objects.values_list("username", "cluster__codice")) == {
        ("m.rossi", "IT"),
        ("m.rossi", "TLC"),
        ("a.bianchi", None),
    }
    di_nuovo = importa_struttura(percorso)
    assert all(v == 0 for v in di_nuovo.conteggi.values()), di_nuovo.conteggi


def test_errore_su_una_riga_blocca_tutto(tmp_path):
    conti = [*CONTI, ("60000009", "Conto (prova)", "ZZ", "01/01/2026", None, None)]
    with pytest.raises(tabelle.FileNonValido, match="Conti, riga 5: cluster inesistente: ZZ"):
        importa_struttura(_completo(tmp_path, conti=conti))
    assert not m.Cluster.objects.filter(codice="XX").exists()
    wbs = [("X.0009.00.000", "Senza cluster (prova)", "Cluster", "", "Sì", "")]
    with pytest.raises(tabelle.FileNonValido, match="una WBS di cluster ha il cluster"):
        importa_struttura(_completo(tmp_path, wbs=wbs))
    with pytest.raises(tabelle.FileNonValido, match="TUTTI è riservato"):
        importa_struttura(_file(tmp_path, cluster=[("TUTTI", "Tutti", "")]))
    with pytest.raises(tabelle.FileNonValido, match="Fogli mancanti"):
        wb = openpyxl.Workbook()
        wb.save(tmp_path / "vuoto.xlsx")
        importa_struttura(tmp_path / "vuoto.xlsx")


def test_responsabili_elenco_completo_o_invariati(tmp_path):
    importa_struttura(_completo(tmp_path))
    r = importa_struttura(_completo(tmp_path, responsabili=[("m.rossi", "IT", "")]))
    assert r.conteggi["assegnazioni_tolte"] == 2
    attive = m.AssegnazioneCluster.objects.filter(archiviato=False)
    assert list(attive.values_list("username", "cluster__codice")) == [("m.rossi", "IT")]
    assert m.AssegnazioneCluster.objects.count() == 3  # le tolte restano, archiviate
    # Foglio vuoto: le assegnazioni restano (es. si aggiornano solo i conti).
    r = importa_struttura(_completo(tmp_path, responsabili=[]))
    assert r.conteggi["assegnazioni_tolte"] == 0
    assert attive.count() == 1
    # Di nuovo nel file: l'assegnazione archiviata torna attiva.
    r = importa_struttura(_completo(tmp_path))
    assert r.conteggi["assegnazioni_create"] == 2 and attive.count() == 3


def test_voci_con_wbs_testuale_collegate_o_segnalate(tmp_path):
    c = f.contratto(fabbisogno=f.fabbisogno(cluster=f.cluster("IT")))
    capex = f.voce(c, natura=m.Natura.CAPEX, codice_wbs_importato="X.0001.00.000")
    opex = f.voce(c, codice_wbs_importato="X.0002.00.000")
    ignota = f.voce(c, natura=m.Natura.CAPEX, codice_wbs_importato="X.9999.99.999")
    r = importa_struttura(_completo(tmp_path))
    capex.refresh_from_db()
    assert capex.wbs.codice_wbs == "X.0001.00.000" and r.conteggi["voci_collegate"] == 1
    codici = {a.codice: a.valore for a in r.anomalie}
    assert codici["WBS_SU_OPEX"] == "X.0002.00.000"
    assert codici["WBS_NON_IN_STRUTTURA"] == "X.9999.99.999"
    opex.refresh_from_db()
    ignota.refresh_from_db()
    assert opex.wbs is None and ignota.wbs is None


def test_conti_senza_cluster_segnalati(tmp_path):
    m.ContoCoge.objects.create(codice="69999999", descrizione="Vecchio", valido_dal=D(2020, 1, 1))
    r = importa_struttura(_completo(tmp_path))
    assert "CONTO_SENZA_CLUSTER" in {a.codice for a in r.anomalie}


def _client(make_user, nome, ruolo):
    make_user(nome, role=ruolo, is_superuser=ruolo == ADMIN, is_staff=ruolo == ADMIN)
    c = Client()
    c.login(username=nome, password=PASSWORD)
    return c


def test_solo_admin_dall_interfaccia(tmp_path, make_user):
    editor = _client(make_user, "u_editor", EDITOR)
    assert "Struttura ICTDINN" not in editor.get(reverse("contratti:importazioni")).content.decode()
    url = reverse("contratti:importazione_carica", args=["struttura"])
    with open(_completo(tmp_path), "rb") as fh:
        file = SimpleUploadedFile("struttura.xlsx", fh.read())
    assert editor.post(url, {"struttura-file": file}).status_code == 403
    admin = _client(make_user, "u_admin", ADMIN)
    risposta = admin.get(reverse("contratti:importazione_modello", args=["struttura"]))
    assert risposta.status_code == 200
    assert (
        'filename="GestioneContratti_ModelloImport_Struttura.xlsx"'
        in risposta["Content-Disposition"]
    )
    file.seek(0)
    risposta = admin.post(url, {"struttura-file": file})
    imp = m.Importazione.objects.get()
    assert risposta.status_code == 302 and imp.stato == m.Importazione.Stato.ANTEPRIMA
    admin.post(reverse("contratti:importazione_conferma", args=[imp.pk]))
    assert m.ContoCoge.objects.filter(clusters__codice="IT").count() == 2


def test_editor_importa_contratti_e_date_solo_nei_propri_cluster(tmp_path, make_user):
    it, tlc = f.cluster("IT"), f.cluster("TLC")
    m.AssegnazioneCluster.objects.create(username="u_editor", cluster=it)
    utente = make_user("u_editor", role=EDITOR)
    tipo = imp_svc.TIPI["contratti"]
    assert [c for c, _ in imp_svc.cluster_importabili(tipo, utente)] == ["IT"]
    file = SimpleUploadedFile("contratti.xlsx", b"")
    with pytest.raises(imp_svc.ImportNonAmmesso, match="solo nei tuoi cluster"):
        imp_svc.carica(tipo, file, {"area": "TLC"}, utente)
    # Date di emissione: i contratti di un altro cluster non si toccano.
    f.contratto(fabbisogno=f.fabbisogno(cluster=it), numero_contratto="5800000101")
    f.contratto(fabbisogno=f.fabbisogno(cluster=tlc), numero_contratto="5800000102")
    csv = "numero_contratto;data_emissione\n5800000101;15/12/2025\n5800000102;15/12/2025\n"
    file = SimpleUploadedFile("date.csv", csv.encode())
    imp = imp_svc.carica(imp_svc.TIPI["date_emissione"], file, {}, utente)
    conteggi = imp.anteprima["conteggi"]
    assert (conteggi["compilati"], conteggi["fuori_cluster"]) == (1, 1)
