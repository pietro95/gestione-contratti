"""Quadratura con il Dashboard Excel (SPEC §7.8).

Per ogni riga del foglio "Contratti" (la tabella su cui lavora il Dashboard) si
ricalcola la formula Excel "Ipotesi di spesa per singolo anno", bug dei mesi compreso,
e la si confronta con la competenza calcolata dall'app sulla voce importata da quella riga.

Formula Excel, per anno: importo / mesi × mesi di calendario toccati nell'anno, estremi
inclusi. Minimo = Corrispettivo Minimo; massimo = Corrispettivo Massimo (formula
minimo + extra, oppure il valore scritto a mano).
"""

import datetime
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
from dateutil.relativedelta import relativedelta

from contratti import adapters
from contratti import models as m
from contratti.engine.date import UN_GIORNO
from contratti.engine.importi import ZERO, arrotonda
from contratti.engine.tipi import Livello
from contratti.importer import normalizza as n
from contratti.importer.carica import chiave_contratto, chiave_voce, raggruppa
from contratti.importer.lettura import leggi, normalizza_riga
from reports.budget import importi_anno

FOGLIO_DASHBOARD = "Contratti"  # il Dashboard somma solo la tabella "Contratti"

CAUSE = {
    "BUG_MESI": "Bug dei mesi: l'Excel conta i mesi di calendario toccati, più della durata.",
    "MESI_VS_GIORNI": "Ripartizione a mesi interi nell'Excel, a mesi frazionati nell'app.",
    "MASSIMO_SOVRASCRITTO": "Corrispettivo Massimo scritto a mano, diverso da minimo + extra.",
    "SCADENZA_SOVRASCRITTA": "Data di scadenza scritta a mano, diversa da inizio + mesi − 1.",
    "MODIFICATA_DOPO_IMPORT": "Importi o date della voce cambiati nell'app dopo l'import.",
    "NON_IMPORTATA": "Riga non trovata nell'app (import non eseguito per quest'area).",
    "ARCHIVIATA": "Voce o contratto archiviati nell'app: fuori dal budget.",
}


def mesi_toccati(inizio: datetime.date, fine: datetime.date) -> int:
    """Mesi di calendario toccati, estremi inclusi (la logica dell'Excel)."""
    return max(0, (fine.year - inizio.year) * 12 + fine.month - inizio.month + 1)


def excel_anno(
    importo: Decimal, mesi: int, inizio: datetime.date, fine: datetime.date, anno: int
) -> Decimal:
    """Quota dell'anno secondo la formula del Dashboard (non arrotondata)."""
    a = max(inizio, datetime.date(anno, 1, 1))
    b = min(fine, datetime.date(anno, 12, 31))
    return importo / Decimal(mesi) * mesi_toccati(a, b)


def scadenza_excel(valore, inizio: datetime.date, mesi: int) -> tuple[datetime.date, bool]:
    """Data di scadenza dell'Excel: EDATE(inizio; mesi) − 1, o il valore scritto a mano.

    Restituisce anche se il valore è scritto a mano.
    """
    calcolata = inizio + relativedelta(months=mesi) - UN_GIORNO
    if isinstance(valore, str) and valore.startswith("="):
        return calcolata, False
    data = n.data(valore)
    return (data, True) if data else (calcolata, False)


@dataclass
class RigaQuadratura:
    foglio: str
    numero: int
    descrizione: str
    fornitore: str
    inizio: datetime.date
    fine: datetime.date
    mesi: int
    excel_minimo: dict[int, Decimal]
    excel_massimo: dict[int, Decimal]
    app_fisso: dict[int, Decimal]
    app_massimo: dict[int, Decimal]
    cause: list[str] = field(default_factory=list)
    contratto_pk: int | None = None

    def differenza(self, anno: int, livello: str = "minimo") -> Decimal:
        if livello == "minimo":
            return self.app_fisso[anno] - arrotonda(self.excel_minimo[anno])
        return self.app_massimo[anno] - arrotonda(self.excel_massimo[anno])


@dataclass
class Quadratura:
    file: str
    area: str
    anni: list[int]
    righe: list[RigaQuadratura]
    escluse: list[tuple[int, str]]  # (riga, motivo): la formula Excel non le conta
    dashboard: dict[int, tuple[Decimal | None, Decimal | None]]  # valori salvati nel file

    def riepilogo(self) -> list[dict]:
        risultato = []
        for anno in self.anni:
            ex_min = arrotonda(sum((r.excel_minimo[anno] for r in self.righe), ZERO))
            ex_max = arrotonda(sum((r.excel_massimo[anno] for r in self.righe), ZERO))
            app_min = sum((r.app_fisso[anno] for r in self.righe), ZERO)
            app_max = sum((r.app_massimo[anno] for r in self.righe), ZERO)
            salvato = self.dashboard.get(anno, (None, None))
            risultato.append(
                {
                    "anno": anno,
                    "excel_minimo": ex_min,
                    "app_fisso": app_min,
                    "diff_minimo": app_min - ex_min,
                    "excel_massimo": ex_max,
                    "app_massimo": app_max,
                    "diff_massimo": app_max - ex_max,
                    "dashboard_minimo": salvato[0],
                    "dashboard_massimo": salvato[1],
                }
            )
        return risultato

    def conteggio_cause(self) -> dict[str, int]:
        conteggio = dict.fromkeys(CAUSE, 0)
        for r in self.righe:
            for c in r.cause:
                conteggio[c] += 1
        return {c: k for c, k in conteggio.items() if k}


def _dashboard(percorso: Path) -> dict[int, tuple[Decimal | None, Decimal | None]]:
    """Valori salvati nel foglio Dashboard (anno, minimo, massimo), se il file li ha."""
    wb = openpyxl.load_workbook(percorso, data_only=True, read_only=True)
    if "Dashboard" not in wb.sheetnames:
        return {}
    valori = {}
    for riga in wb["Dashboard"].iter_rows(min_col=6, max_col=8, values_only=True):
        anno, minimo, massimo = (list(riga) + [None, None, None])[:3]
        if isinstance(anno, int) and 2000 <= anno <= 2100:
            valori[anno] = tuple(
                Decimal(str(v)) if isinstance(v, int | float) else None for v in (minimo, massimo)
            )
    wb.close()
    return valori


def _app(voce: m.VoceContratto | None, anni: list[int]):
    """Competenza dell'app per anno (fisso e massimo) e cause legate allo stato nell'app."""
    vuoto = dict.fromkeys(anni, ZERO)
    if voce is None:
        return dict(vuoto), dict(vuoto), ["NON_IMPORTATA"]
    contratto = voce.contratto
    if voce.archiviato or contratto.archiviato or contratto.fabbisogno.archiviato:
        return dict(vuoto), dict(vuoto), ["ARCHIVIATA"]
    voci_engine, _ = adapters.voci_contratto(contratto)
    ve = next((v for v in voci_engine if v.riferimento == voce.pk), None)
    if ve is None:
        return dict(vuoto), dict(vuoto), ["ARCHIVIATA"]
    fisso, massimo = dict(vuoto), dict(vuoto)
    for anno in anni:
        importi, _ = importi_anno(ve, anno)
        fisso[anno] = importi["competenza"][Livello.FISSO]
        massimo[anno] = importi["competenza"][Livello.MASSIMO]
    return fisso, massimo, []


def quadra(
    percorso: str | Path, *, area: str, anni: list[int] | None = None, oggi=None
) -> Quadratura:
    percorso = Path(percorso)
    libro = leggi(percorso)
    grezze = {(g.foglio, g.numero): g for g in libro.righe}
    righe = [normalizza_riga(g) for g in libro.righe]
    dashboard = _dashboard(percorso)
    if anni is None:
        oggi = oggi or datetime.date.today()
        anni = sorted(dashboard) or list(range(oggi.year - 2, oggi.year + 3))

    chiavi: dict[tuple[str, int], str] = {}
    for chiave_gruppo, gruppo in raggruppa(righe).items():
        chiave = chiave_contratto(area, chiave_gruppo)
        for i, r in enumerate(gruppo):
            chiavi[(r.foglio, r.numero)] = chiave_voce(chiave, i)
    voci = m.VoceContratto.objects.select_related("contratto__fabbisogno").in_bulk(
        list(chiavi.values()), field_name="chiave_import"
    )

    risultato: list[RigaQuadratura] = []
    escluse: list[tuple[int, str]] = []
    for r in righe:
        if r.foglio != FOGLIO_DASHBOARD:
            continue
        g = grezze[(r.foglio, r.numero)]
        mesi_excel = n.intero(g.valori.get("mesi"))
        if r.data_inizio is None:
            escluse.append((r.numero, 'Data di attivazione mancante ("n/a" nell\'Excel).'))
            continue
        if not mesi_excel:
            escluse.append((r.numero, "Mesi contratto mancanti: la formula Excel va in errore."))
            continue
        fine, scadenza_a_mano = scadenza_excel(g.valori.get("scadenza"), r.data_inizio, mesi_excel)
        minimo = r.minimo or ZERO
        if g.massimo_formula:
            massimo = minimo + (r.extra or ZERO)
        else:
            massimo, _ = n.importo(g.valori.get("massimo"), "massimo")
            massimo = massimo or ZERO
        voce = voci.get(chiavi[(r.foglio, r.numero)])
        app_fisso, app_massimo, cause = _app(voce, anni)
        riga = RigaQuadratura(
            foglio=r.foglio,
            numero=r.numero,
            descrizione=r.descrizione or "",
            fornitore=r.fornitore or "",
            inizio=r.data_inizio,
            fine=fine,
            mesi=mesi_excel,
            excel_minimo={a: excel_anno(minimo, mesi_excel, r.data_inizio, fine, a) for a in anni},
            excel_massimo={
                a: excel_anno(massimo, mesi_excel, r.data_inizio, fine, a) for a in anni
            },
            app_fisso=app_fisso,
            app_massimo=app_massimo,
            cause=cause,
            contratto_pk=voce.contratto_id if voce else None,
        )
        if mesi_toccati(r.data_inizio, fine) != mesi_excel:
            riga.cause.append("BUG_MESI")
        if not g.massimo_formula and massimo != minimo + (r.extra or ZERO):
            riga.cause.append("MASSIMO_SOVRASCRITTO")
        if scadenza_a_mano and fine != r.data_inizio + relativedelta(months=mesi_excel) - UN_GIORNO:
            riga.cause.append("SCADENZA_SOVRASCRITTA")
        if voce is not None and not cause and _modificata(voce, r):
            riga.cause.append("MODIFICATA_DOPO_IMPORT")
        differenze = any(riga.differenza(a, liv) for a in anni for liv in ("minimo", "massimo"))
        if differenze and not riga.cause:
            riga.cause.append("MESI_VS_GIORNI")
        risultato.append(riga)
    return Quadratura(percorso.name, area, anni, risultato, escluse, dashboard)


def _modificata(voce: m.VoceContratto, r) -> bool:
    c = voce.contratto
    return (
        voce.importo_fisso != (r.minimo or ZERO)
        or voce.massimale_variabile != (r.extra or ZERO)
        or c.data_inizio != r.data_inizio
        or c.durata_mesi != r.mesi
    )
