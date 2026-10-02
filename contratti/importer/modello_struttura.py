# ruff: noqa: E501 — testi e formule del modello approvato, riportati senza modifiche.
"""Modello dell'import della struttura ICTDINN (Rev00, approvato da Pietro il 02/10/2026).

Si genera dal codice, non si versiona come file: gli .xlsx fuori da ``tests/fixtures`` sono
esclusi dal repository (dati reali). Contiene solo istruzioni, esempi fittizi e i codici dei
cluster proposti.
"""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NOME = "2026.10.02_PFE_Modello_import_struttura_ICTDINN_Rev00.xlsx"
F = "Arial"
BLU = "003F6B"
GIALLO = PatternFill("solid", fgColor="FFF2CC")
BIANCO_H = PatternFill("solid", fgColor="DCE6F0")
TIT = Font(name=F, bold=True, size=14, color=BLU)
H = Font(name=F, bold=True, size=10, color="000000")
N = Font(name=F, size=10)
NI = Font(name=F, size=10, italic=True, color="555555")
B = Font(name=F, bold=True, size=10)
SOTTILE = Side(style="thin", color="B7C3D0")
BORDO = Border(left=SOTTILE, right=SOTTILE, top=SOTTILE, bottom=SOTTILE)
RIGHE = 500  # righe predisposte per validazioni e controlli


def crea() -> Workbook:
    """Il modello Rev00: foglio Istruzioni, poi Cluster, Responsabili, Conti e WBS."""
    wb = Workbook()

    # --- Istruzioni ----------------------------------------------------------------
    ist = wb.active
    ist.title = "Istruzioni"
    ist.sheet_view.showGridLines = False
    ist.column_dimensions["A"].width = 3
    for c, w in zip("BCDEFG", (22, 28, 18, 16, 16, 24), strict=False):
        ist.column_dimensions[c].width = w
    r = 2
    ist.cell(r, 2, "Modello di import — struttura dell'ufficio ICTDINN").font = TIT
    r += 1
    ist.cell(
        r,
        2,
        "Cluster, responsabili, piano dei conti e WBS per il Registro contratti (fase F9). Rev00 — 02/10/2026.",
    ).font = NI
    r += 2
    testi = [
        ("Come si usa", None),
        (
            "1.",
            "Compila i fogli Cluster, Responsabili, Conti e WBS. Non rinominare i fogli e non cambiare le intestazioni della riga 1.",
        ),
        (
            "2.",
            "Intestazione gialla = colonna obbligatoria; azzurra = facoltativa. Il foglio Istruzioni e la colonna «Controllo» non vengono importati.",
        ),
        (
            "3.",
            "La colonna «Controllo» segnala gli errori mentre scrivi (codici ripetuti, cluster inesistenti, dati mancanti). Deve essere vuota su tutte le righe.",
        ),
        (
            "4.",
            "Codici come testo: conti e WBS restano come li scrivi (zeri iniziali compresi). Date nel formato gg/mm/aaaa.",
        ),
        (
            "5.",
            "Carica il file da «Gestione → Import → Struttura ICTDINN»: prima vedi l'anteprima con le anomalie, poi confermi. Rieseguire l'import con lo stesso file non cambia nulla.",
        ),
        (
            "6.",
            "Nessuna cancellazione: un conto o una WBS tolti dal file restano nel registro. Per dismetterli compila «valido_al» (conti) o metti «No» in «attiva» (WBS).",
        ),
        ("", None),
        ("Regole per foglio", None),
        (
            "Cluster",
            "Un cluster è un'area dell'ufficio che raggruppa più conti (es. IT, TLC, OT). Il codice è breve e non cambia nel tempo: lo usano gli altri fogli.",
        ),
        (
            "Responsabili",
            "Una riga per ogni coppia utente–cluster. Utente = nome utente di dominio (sAMAccountName), anche se non è ancora entrato nell'app. Cluster = codice, oppure TUTTI (es. addetto amministrativo). L'Editor modifica solo i fabbisogni dei suoi cluster; il ruolo (Admin, Editor, Reader) resta quello dei gruppi AD.",
        ),
        (
            "Conti",
            "Ogni conto appartiene a un solo cluster. La coppia codice + valido_dal identifica il conto: con un valido_dal diverso è un conto nuovo (validità temporale). «sostituito_da» = codice del conto che lo sostituisce dalla fine della validità.",
        ),
        (
            "WBS",
            "Tipo «Cluster»: CAPEX generico del cluster, entra nel suo budget CAPEX (cluster obbligatorio). Tipo «Progetto»: WBS a progetto (es. nuovi dissalatori), fuori dai budget dei cluster, i contratti restano registrati e consultabili per WBS; i contratti finanziati stanno sempre su una WBS a progetto.",
        ),
        ("", None),
        ("Esempi (valori fittizi, solo per il formato: non copiarli nei fogli)", None),
    ]
    for a, b in testi:
        if b is None:
            if a:
                ist.cell(r, 2, a).font = Font(name=F, bold=True, size=11, color=BLU)
            r += 1
            continue
        ist.cell(r, 2, a).font = B
        c = ist.cell(r, 3, b)
        c.font = N
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ist.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        ist.row_dimensions[r].height = 15 * max(1, len(b) // 95 + 1)
        ist.cell(r, 2).alignment = Alignment(vertical="top")
        r += 1

    esempi = [
        ("Cluster", ["codice", "nome", "note"], ["XX", "Cluster di esempio", ""]),
        ("Responsabili", ["utente", "cluster", "note"], ["m.rossi", "XX", "Responsabile"]),
        (
            "Conti",
            ["codice", "descrizione", "cluster", "valido_dal", "valido_al", "sostituito_da"],
            ["60000000", "Canoni software (esempio)", "XX", "01/01/2026", "", ""],
        ),
        (
            "WBS",
            ["codice_wbs", "descrizione", "tipo", "cluster", "attiva", "note"],
            ["X.0000.00.000", "Investimenti di esempio", "Cluster", "XX", "Sì", ""],
        ),
    ]
    for nome, intest, riga in esempi:
        ist.cell(r, 2, nome).font = B
        r += 1
        for j, (h, v) in enumerate(zip(intest, riga, strict=False)):
            c = ist.cell(r, 2 + j, h)
            c.font = H
            c.fill = BIANCO_H
            c.border = BORDO
            c = ist.cell(r + 1, 2 + j, v)
            c.font = NI
            c.border = BORDO
        r += 3

    # --- Fogli dati ----------------------------------------------------------------
    def foglio(titolo, colonne, larghezze, obbligatorie, testo=()):
        ws = wb.create_sheet(titolo)
        ws.freeze_panes = "A2"
        for j, (nome, w) in enumerate(zip(colonne, larghezze, strict=False), start=1):
            c = ws.cell(1, j, nome)
            c.font = H
            c.fill = GIALLO if nome in obbligatorie else BIANCO_H
            c.border = BORDO
            c.alignment = Alignment(vertical="center")
            ws.column_dimensions[c.column_letter].width = w
        ws.row_dimensions[1].height = 20
        for j, nome in enumerate(colonne, start=1):
            if nome in testo:
                for i in range(2, RIGHE + 2):
                    ws.cell(i, j).number_format = "@"
        for i in range(2, RIGHE + 2):
            for j in range(1, len(colonne) + 1):
                ws.cell(i, j).font = N
        return ws

    def controllo(ws, col, formula):
        """Colonna Controllo: testo dell'errore, vuoto se la riga è a posto."""
        c = ws.cell(1, col, "Controllo")
        c.font = Font(name=F, bold=True, size=10, color="9B1C13")
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        c.border = BORDO
        ws.column_dimensions[c.column_letter].width = 42
        for i in range(2, RIGHE + 2):
            cella = ws.cell(i, col, formula.format(r=i))
            cella.font = Font(name=F, size=10, color="9B1C13")

    def tendina(ws, rif, colonna, titolo, messaggio):
        dv = DataValidation(
            type="list",
            formula1=rif,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle=titolo,
            error=messaggio,
        )
        ws.add_data_validation(dv)
        dv.add(f"{colonna}2:{colonna}{RIGHE + 1}")

    CLUSTER = "Cluster!$A$2:$A$100"

    cl = foglio("Cluster", ["codice", "nome", "note"], [12, 32, 40], {"codice", "nome"}, {"codice"})
    # Cluster indicati da Pietro (02/10/2026): nomi organizzativi, da completare o correggere.
    for i, (cod, nome) in enumerate(
        [
            ("IT", "IT"),
            ("TLC", "TLC"),
            ("OT", "OT"),
            ("SICFIS", "Sicurezza fisica"),
            ("INN", "Innovazione"),
        ],
        start=2,
    ):
        cl.cell(i, 1, cod).font = Font(name=F, size=10, color="0000FF")
        cl.cell(i, 2, nome).font = Font(name=F, size=10, color="0000FF")
    cl.cell(2, 3, "Codici proposti: verificali prima di caricare.").font = NI
    controllo(
        cl,
        4,
        '=IF(AND(A{r}="",B{r}=""),"",IF(A{r}="","Manca il codice",IF(B{r}="","Manca il nome",'
        'IF(COUNTIF($A$2:$A$100,A{r})>1,"Codice ripetuto",IF(UPPER(A{r})="TUTTI","TUTTI è riservato","")))))',
    )

    re_ = foglio(
        "Responsabili",
        ["utente", "cluster", "note"],
        [22, 14, 40],
        {"utente", "cluster"},
        {"utente"},
    )
    dv = DataValidation(
        type="custom",
        formula1=f'=OR(B2="TUTTI",COUNTIF({CLUSTER},B2)>0)',
        showErrorMessage=True,
        errorTitle="Cluster",
        error="Scegli un codice del foglio Cluster, oppure TUTTI.",
    )
    re_.add_data_validation(dv)
    dv.add(f"B2:B{RIGHE + 1}")
    controllo(
        re_,
        4,
        '=IF(AND(A{r}="",B{r}=""),"",IF(A{r}="","Manca l\'utente",IF(B{r}="","Manca il cluster",'
        'IF(AND(B{r}<>"TUTTI",COUNTIF(' + CLUSTER + ',B{r})=0),"Cluster inesistente",'
        "IF(COUNTIFS($A$2:$A$"
        + str(RIGHE + 1)
        + ",A{r},$B$2:$B$"
        + str(RIGHE + 1)
        + ',B{r})>1,"Riga ripetuta","")))))',
    )

    co = foglio(
        "Conti",
        ["codice", "descrizione", "cluster", "valido_dal", "valido_al", "sostituito_da"],
        [14, 40, 12, 13, 13, 15],
        {"codice", "descrizione", "cluster", "valido_dal"},
        {"codice", "sostituito_da"},
    )
    for i in range(2, RIGHE + 2):
        for col in (4, 5):
            co.cell(i, col).number_format = "DD/MM/YYYY"
    tendina(co, f"={CLUSTER}", "C", "Cluster", "Scegli un codice del foglio Cluster.")
    dv = DataValidation(
        type="date",
        operator="greaterThan",
        formula1="36526",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Data",
        error="Scrivi una data gg/mm/aaaa.",
    )
    co.add_data_validation(dv)
    dv.add(f"D2:E{RIGHE + 1}")
    n = RIGHE + 1
    controllo(
        co,
        7,
        '=IF(AND(A{r}="",B{r}=""),"",IF(A{r}="","Manca il codice",IF(B{r}="","Manca la descrizione",'
        'IF(C{r}="","Manca il cluster",IF(COUNTIF(' + CLUSTER + ',C{r})=0,"Cluster inesistente",'
        'IF(D{r}="","Manca valido_dal",IF(AND(E{r}<>"",E{r}<D{r}),"valido_al prima di valido_dal",'
        "IF(COUNTIFS($A$2:$A$"
        + str(n)
        + ",A{r},$D$2:$D$"
        + str(n)
        + ',D{r})>1,"Codice e valido_dal ripetuti",'
        'IF(AND(F{r}<>"",COUNTIF($A$2:$A$' + str(n) + ',F{r})=0),"sostituito_da non è nel foglio",'
        'IF(F{r}=A{r},"Un conto non sostituisce se stesso",""))))))))))',
    )

    wbs = foglio(
        "WBS",
        ["codice_wbs", "descrizione", "tipo", "cluster", "attiva", "note"],
        [18, 40, 11, 12, 9, 36],
        {"codice_wbs", "descrizione", "tipo"},
        {"codice_wbs"},
    )
    tendina(wbs, '"Cluster,Progetto"', "C", "Tipo", "Scegli Cluster o Progetto.")
    tendina(wbs, f"={CLUSTER}", "D", "Cluster", "Scegli un codice del foglio Cluster.")
    tendina(wbs, '"Sì,No"', "E", "Attiva", "Sì o No (vuoto = Sì).")
    wbs.cell(1, 4).comment = Comment(
        "Obbligatorio per il tipo Cluster. Per il tipo Progetto è facoltativo: indica chi segue il progetto, ma la WBS resta fuori dal budget del cluster.",
        "Registro",
    )
    n = RIGHE + 1
    controllo(
        wbs,
        7,
        '=IF(AND(A{r}="",B{r}=""),"",IF(A{r}="","Manca il codice WBS",IF(B{r}="","Manca la descrizione",'
        'IF(C{r}="","Manca il tipo",IF(AND(C{r}="Cluster",D{r}=""),"Tipo Cluster: manca il cluster",'
        'IF(AND(D{r}<>"",COUNTIF(' + CLUSTER + ',D{r})=0),"Cluster inesistente",'
        "IF(COUNTIF($A$2:$A$" + str(n) + ',A{r})>1,"Codice WBS ripetuto","")))))))',
    )

    for ws in wb.worksheets[1:]:
        ws.sheet_properties.tabColor = BLU
    ist.sheet_properties.tabColor = "177F42"
    # Stampa / PDF: istruzioni intere, dei fogli dati solo le intestazioni e le prime righe.
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if ws.title != "Istruzioni":
            ws.print_area = f"A1:{ws.cell(1, ws.max_column).column_letter}12"
    return wb
