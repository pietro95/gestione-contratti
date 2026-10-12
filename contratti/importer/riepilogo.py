"""Import una tantum del «Riepilogo RDA Ordini» (SPEC §2.3, §7; decisioni di Pietro del
07/10/2026).

- Solo le RDA con data dal 01/04/2025 (numerazione SAP) e con un conto di un cluster del
  registro (ICTDINN); le altre vanno nel report (``FUORI_PERIMETRO``).
- L'intestazione è la riga con «Numero» e «Data…» nelle prime due colonne; la riga sopra ha i
  gruppi (RDA, contratto, fattura, DDT, attestazione, …). Le celle unite si ricompongono: ogni
  riga eredita il valore della cella unita che la copre. Una RDA occupa le righe dal suo
  numero alla riga prima del numero successivo.
- Le celle con più valori si dividono per riga (a capo); stati, date e importi si normalizzano
  e ogni sostituzione si segnala.
- Centro di costo e conto invertiti (il «centro di costo» è un conto del piano dei conti e il
  «conto» no) si scambiano e si segnalano (``CONTO_CDC_SCAMBIATI``).
- Contratto già registrato (numero SAP) → si collega alla RDA; altrimenti si crea in un
  fabbisogno per cluster e categoria «Da classificare» (``CONTRATTO_CREATO``), una tantum se
  non ha inizio e fine delle prestazioni.
- Fatture, DDT, attestazioni e numeri SAP (ODA/RDC, MAP, MIGO) vanno al contratto della loro
  riga o, se la riga non ne ha, all'ultimo contratto sopra nella stessa RDA. Le fatture sono
  storiche: valore base = imponibile, convalidate, trasferite alla data della fattura se la nota
  dice «trasferita».
- «Chiuso» diventa la chiusura dei contratti della RDA alla data dell'import; senza
  attestazione si segnala (``CHIUSO_SENZA_ATTESTAZIONE``). Il basket del masterplan non si
  importa. Ogni RDA ha una sola posizione; il referente resta vuoto (``REFERENTE_MANCANTE``).
- Idempotente: una RDA già registrata con lo stesso numero non si reimporta.
"""

import datetime
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.db import transaction
from django.utils import timezone

from contratti import models as m
from contratti.importer import normalizza, tabelle
from contratti.importer.report import Report
from contratti.importer.tabelle import FileNonValido

INIZIO_SAP = datetime.date(2025, 4, 1)
RE_RDA = re.compile(r"^51\d{8}$")
RE_CONTRATTO = re.compile(r"^\d{10}$")
NULLI = {"", "-", "--", "---", "n/a", "n/d", "na", "nd", "non previsto in sap", "nessuno"}
EPOCA_EXCEL = datetime.date(1899, 12, 30)
CATEGORIA = "Da classificare"
PROCEDURA = "da_verificare"
FOGLIO = "Riepilogo"

# (gruppo, sottotitolo) → campo. Il gruppo è la cella unita della riga sopra l'intestazione.
COLONNE = {
    ("richiesta d'acquisto", "numero"): "rda",
    ("richiesta d'acquisto", "data"): "data_rda",
    ("richiesta d'acquisto", "descrizione"): "descrizione",
    ("richiesta d'acquisto", "wbs"): "wbs",
    ("richiesta d'acquisto", "centro di costo"): "centro_costo",
    ("richiesta d'acquisto", "codice conto"): "conto",
    ("richiesta d'acquisto", "richiedente"): "richiedente",
    ("richiesta d'acquisto", "rup"): "rup",
    ("richiesta d'acquisto", "cup"): "cup",
    ("contratto", "codifica"): "codifica",
    ("contratto", "numerazione sap"): "contratto",
    ("contratto", "fornitore"): "fornitore",
    ("contratto", "data inizio"): "inizio",
    ("contratto", "data fine"): "fine",
    ("contratto", "data"): "data_contratto",
    ("contratto", "importo base"): "importo_base",
    ("contratto", "importo massimo"): "importo_massimo",
    ("fattura", "numero"): "fattura",
    ("fattura", "data"): "data_fattura",
    ("fattura", "imponibile"): "imponibile",
    ("ddt", "numero"): "ddt",
    ("ddt", "data"): "data_ddt",
    ("attestazione", "numero"): "attestazione",
    ("attestazione", "data"): "data_attestazione",
    ("percorso", ""): "percorso",
    ("stato", ""): "stato",
    ("note", ""): "note",
    ("oda", ""): "oda",
    ("map", ""): "map",
    ("migo", ""): "migo",
}


class _Rollback(Exception):
    pass


# --- Normalizzazione -------------------------------------------------------------------------


def pulisci(valore) -> str | None:
    """Testo senza invisibili, apici e virgolette ai lati; ``None`` per i nulli."""
    if valore is None:
        return None
    if isinstance(valore, float) and valore.is_integer():
        valore = int(valore)
    s = normalizza.testo(valore)
    if s is None:
        return None
    s = s.strip("'\"‘’“” ").strip()
    return None if s.lower() in NULLI else s


def valori(valore) -> list[str]:
    """I valori di una cella con più righe (a capo), senza i nulli."""
    if valore is None:
        return []
    if not isinstance(valore, str):
        v = pulisci(valore)
        return [v] if v else []
    return [v for parte in valore.splitlines() if (v := pulisci(parte))]


def parti(valore) -> list:
    """Come ``valori``, ma lascia date e numeri com'erano (per ``data`` e ``importo``)."""
    if valore is None:
        return []
    if isinstance(valore, str):
        return valori(valore)
    return [valore]


def data(valore) -> datetime.date | None:
    if isinstance(valore, datetime.datetime):
        return valore.date()
    if isinstance(valore, datetime.date):
        return valore
    if isinstance(valore, int | float) and 20000 < valore < 80000:
        return EPOCA_EXCEL + datetime.timedelta(days=int(valore))
    s = pulisci(valore)
    if s is None:
        return None
    for formato in ("%d/%m/%Y", "%d/%m/%y", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return datetime.datetime.strptime(s, formato).date()
        except ValueError:
            continue
    return None


def importo(valore) -> Decimal | None:
    if isinstance(valore, str):
        valore = pulisci(valore)
        if valore is None:
            return None
    d, problemi = normalizza.importo(valore, "importo")
    return None if problemi else d


def stato(valore) -> str:
    s = (pulisci(valore) or "").lower()
    if "chius" in s:
        return "chiuso"
    if "apert" in s:
        return "aperto"
    return ""


def codice_numerico(valore) -> str | None:
    """«999901 (Rossi)» → «999901»; numeri → testo."""
    s = pulisci(valore)
    if s is None:
        return None
    trovato = re.match(r"^\s*(\d+)", s)
    return trovato.group(1) if trovato else s


# --- Lettura del foglio ----------------------------------------------------------------------


def _chiave_gruppo(testo: str) -> str:
    t = (testo or "").lower()
    for chiave in (
        "richiesta d'acquisto",
        "rpa",
        "contratto",
        "fattura",
        "ddt",
        "attestazione",
        "percorso",
        "stato",
        "basket",
        "note",
    ):
        if chiave in t:
            return chiave
    for chiave in ("oda", "map", "migo"):
        if re.search(rf"\b{chiave}\b", t):
            return chiave
    return t


def _chiave_sotto(testo: str) -> str:
    t = " ".join((testo or "").lower().split())
    for chiave in (
        "data inizio",
        "data fine",
        "numerazione sap",
        "codifica",
        "importo base",
        "importo massimo",
        "centro di costo",
        "codice conto",
        "richiedente",
        "wbs",
        "descrizione",
        "fornitore",
        "imponibile",
        "rup",
        "cup",
        "numero",
    ):
        if chiave in t:
            return chiave
    return "data" if t.startswith("data") else t


@dataclass
class Foglio:
    nome: str
    righe: list[dict] = field(default_factory=list)  # numero di riga Excel in "_riga"


def leggi(percorso: Path) -> Foglio:
    try:
        wb = openpyxl.load_workbook(percorso, data_only=True)
    except Exception as exc:
        raise FileNonValido(f"Il file non è un Excel leggibile ({exc}).") from exc
    ws = wb.worksheets[0]
    intestazione = None
    for r in range(1, 20):
        a = (str(ws.cell(r, 1).value or "")).strip().lower()
        b = (str(ws.cell(r, 2).value or "")).strip().lower()
        if a == "numero" and b.startswith("data"):
            intestazione = r
            break
    if intestazione is None:
        raise FileNonValido("Intestazione non trovata: serve la riga «Numero», «Data…».")
    ultima = 0
    for c in range(1, 80):
        if ws.cell(intestazione, c).value or ws.cell(intestazione - 1, c).value:
            ultima = c
    colonne = {}
    gruppo = ""
    for c in range(1, ultima + 1):
        sopra = ws.cell(intestazione - 1, c).value
        if sopra:
            gruppo = _chiave_gruppo(str(sopra))
        sotto = ws.cell(intestazione, c).value
        chiave = (gruppo, _chiave_sotto(str(sotto)) if sotto else "")
        if chiave in COLONNE and COLONNE[chiave] not in colonne.values():
            colonne[c] = COLONNE[chiave]
    mancanti = {"rda", "data_rda", "conto", "contratto", "fattura"} - set(colonne.values())
    if mancanti:
        raise FileNonValido(f"Colonne non trovate: {', '.join(sorted(mancanti))}.")
    # Celle unite: ogni cella della zona eredita il valore in alto a sinistra.
    unite = {}
    for zona in ws.merged_cells.ranges:
        if zona.max_row <= intestazione:
            continue
        valore = ws.cell(zona.min_row, zona.min_col).value
        for r in range(max(zona.min_row, intestazione + 1), zona.max_row + 1):
            for c in range(zona.min_col, min(zona.max_col, ultima) + 1):
                unite[(r, c)] = valore
    foglio = Foglio(ws.title)
    for r in range(intestazione + 1, ws.max_row + 1):
        riga = {"_riga": r}
        for c, campo in colonne.items():
            riga[campo] = unite.get((r, c), ws.cell(r, c).value)
        if any(v not in (None, "") for k, v in riga.items() if k != "_riga"):
            foglio.righe.append(riga)
    return foglio


def blocchi(foglio: Foglio) -> list[list[dict]]:
    """Righe raggruppate per RDA: una RDA va dal suo numero alla riga prima del successivo."""
    risultato: list[list[dict]] = []
    precedente = None
    for riga in foglio.righe:
        numero = pulisci(riga["rda"])
        if numero is not None and numero != precedente:
            risultato.append([riga])
            precedente = numero
        elif risultato:
            risultato[-1].append(riga)
    return risultato


# --- Import -----------------------------------------------------------------------------------


class Importatore:
    def __init__(self, report: Report, oggi: datetime.date):
        self.report = report
        self.oggi = oggi
        self.conti = {}
        for c in m.ContoCoge.objects.select_related("cluster").order_by("valido_dal"):
            self.conti[c.codice] = c  # l'ultimo per validità
        self.fornitori = {}
        for f in m.Fornitore.objects.filter(archiviato=False):
            for nome in (f.ragione_sociale, f.nome_breve):
                if k := normalizza.chiave_nome(nome):
                    self.fornitori.setdefault(k, f)
        self.wbs = {w.codice_wbs.strip().upper(): w for w in m.WBS.objects.all()}
        self.richiedenti = {r.codice: r for r in m.RichiedenteSAP.objects.all()}
        self.procedura = (
            m.Procedura.objects.filter(codice=PROCEDURA).first() or m.Procedura.objects.first()
        )
        self.categoria = m.Categoria.objects.filter(nome=CATEGORIA).first()
        if self.categoria is None:
            self.categoria = m.Categoria.objects.create(nome=CATEGORIA)

    def anomalia(self, riga, codice, campo="", valore="", messaggio=""):
        self.report.anomalia(FOGLIO, riga, codice, campo, valore or "", messaggio)

    def conta(self, chiave, n=1):
        self.report.conteggi[chiave] += n

    # --- RDA ---------------------------------------------------------------------------

    def importa_blocco(self, blocco: list[dict]) -> None:
        primo = blocco[0]
        n = primo["_riga"]
        numero = re.sub(r"\D", "", pulisci(primo["rda"]) or "")
        if not numero:
            return  # legenda o note in fondo al foglio: non è una RDA
        self.conta("rda_lette")
        data_rda = data(primo["data_rda"])
        prima = len(self.report.anomalie)
        conto, centro = self._conto_e_centro(primo, n, data_rda)
        if tabelle.centro_costo_da_segnalare(centro):
            self.anomalia(n, "CENTRO_COSTO_NON_IN_TABELLA", "centro_costo", centro,
                          tabelle.MESSAGGIO_CENTRO_COSTO)  # fmt: skip
        if data_rda is None or data_rda < INIZIO_SAP or conto is None or not conto.cluster_id:
            # Fuori perimetro: solo questa anomalia, non quelle del conto.
            del self.report.anomalie[prima:]
        if data_rda is None or data_rda < INIZIO_SAP:
            self.conta("fuori_perimetro")
            self.anomalia(
                n,
                "FUORI_PERIMETRO",
                "data",
                pulisci(primo["data_rda"]),
                "RDA prima del 01/04/2025 (numerazione SAP) o senza data.",
            )
            return
        if conto is None or conto.cluster_id is None:
            self.conta("fuori_perimetro")
            self.anomalia(
                n,
                "FUORI_PERIMETRO",
                "conto",
                pulisci(primo["conto"]),
                "Conto non di un cluster del registro: non è ICTDINN.",
            )
            return
        if not RE_RDA.match(numero):
            self.conta("non_importate")
            self.anomalia(
                n,
                "NUMERO_RDA_NON_VALIDO",
                "numero",
                pulisci(primo["rda"]),
                "Il numero RDA SAP ha 10 cifre e inizia con 51.",
            )
            return
        if m.RichiestaAcquisto.objects.filter(numero_rda=numero, archiviato=False).exists():
            self.conta("rda_gia_presenti")
            return
        rda = self._rda(primo, numero, data_rda, conto, centro)
        self.anomalia(
            n,
            "REFERENTE_MANCANTE",
            "referente",
            numero,
            "RDA importata senza referente: si assegna dalla pagina della RDA.",
        )
        contratti = self._contratti_e_documenti(blocco, rda, conto, centro)
        annullata = any(
            "annullat" in (pulisci(r.get("codifica")) or "").lower()
            or "annullat" in (pulisci(r.get("contratto")) or "").lower()
            for r in blocco
        )
        if annullata and not contratti:
            m.EventoRDA.objects.create(
                rda=rda,
                tipo=m.EventoRDA.Tipo.ANNULLAMENTO,
                data=data_rda,
                nota="Annullata secondo il Riepilogo RDA Ordini.",
            )
            self.conta("rda_annullate")
        if any(stato(r.get("stato")) == "chiuso" for r in blocco):
            for c in contratti:
                self._chiudi(c, n)

    def _conti_in(self, valore) -> list[m.ContoCoge]:
        """I conti del piano dei conti scritti nella cella (anche più d'uno, con altro testo:
        «G1517 611021», «611111 (2025) 611138 (dal 2026)»)."""
        trovati = []
        for codice in re.findall(r"\d{5,10}", " ".join(valori(valore))):
            conto = self.conti.get(codice)
            if conto is not None and conto not in trovati:
                trovati.append(conto)
        return trovati

    @staticmethod
    def _valido(conti: list[m.ContoCoge], alla_data) -> m.ContoCoge:
        for c in conti:
            if c.valido_dal <= alla_data and (c.valido_al is None or alla_data <= c.valido_al):
                return c
        return conti[0]

    def _conto_e_centro(self, riga, n, alla_data=None):
        alla_data = alla_data or self.oggi
        testo_conto = " ".join(valori(riga.get("conto")))
        testo_centro = " ".join(valori(riga.get("centro_costo")))
        conti = self._conti_in(riga.get("conto"))
        centro = testo_centro
        if not conti and (scambiati := self._conti_in(riga.get("centro_costo"))):
            self.anomalia(
                n,
                "CONTO_CDC_SCAMBIATI",
                "conto",
                f"{testo_conto} / {testo_centro}",
                "Centro di costo e conto invertiti nel file: scambiati.",
            )
            conti, centro = scambiati, testo_conto
        if not conti:
            return None, centro[:20]
        conto = self._valido(conti, alla_data)
        if len(conti) > 1:
            self.anomalia(
                n,
                "PIU_VALORI",
                "conto",
                testo_conto or testo_centro,
                f"Più conti nella cella: si usa {conto.codice}, valido alla data della RDA.",
            )
        # Il centro di costo: il primo codice che non è il conto.
        codici = [c for c in re.findall(r"[0-9A-Z]{4,10}", centro) if c != conto.codice]
        return conto, (codici[0] if codici else centro)[:20]

    def _rda(self, primo, numero, data_rda, conto, centro) -> m.RichiestaAcquisto:
        descrizione = (pulisci(primo.get("descrizione")) or "RDA importata dal Riepilogo")[:250]
        richiedente_testo = pulisci(primo.get("richiedente")) or ""
        richiedente = self.richiedenti.get(codice_numerico(primo.get("richiedente")) or "")
        note = ["Importata dal Riepilogo RDA Ordini."]
        if richiedente_testo and richiedente is None:
            note.append(f"Gestore/Richiedente SAP nel file: {richiedente_testo}.")
        if nota := pulisci(primo.get("note")):
            note.append(nota)
        rda = m.RichiestaAcquisto(
            origine=m.RichiestaAcquisto.Origine.IMPORT,
            cluster=conto.cluster,
            descrizione=descrizione,
            numero_rda=numero,
            data_creazione=data_rda,
            data_inserimento=data_rda,
            data_rilascio=data_rda,
            richiedente=richiedente,
            rup=(pulisci(primo.get("rup")) or "")[:150],
            cup=(pulisci(primo.get("cup")) or "")[:20],
            percorso_file_server=(pulisci(primo.get("percorso")) or "")[:500],
            note="\n".join(note),
        )
        rda._change_reason = "Import del Riepilogo RDA Ordini"
        rda.save()
        wbs = self._wbs(primo)
        base = next((b for r in [primo] if (b := importo(r.get("importo_base"))) is not None), None)
        m.PosizioneRDA.objects.create(
            rda=rda,
            numero=1,
            descrizione_breve=descrizione,
            importo_unitario=base or Decimal("0.00"),
            natura=m.Natura.CAPEX if wbs else m.Natura.OPEX,
            conto=conto,
            centro_costo=centro,
            wbs=wbs,
        )
        self.conta("rda_create")
        return rda

    def _wbs(self, riga):
        codice = pulisci(riga.get("wbs"))
        return self.wbs.get(codice.upper()) if codice else None

    # --- Contratti ----------------------------------------------------------------------

    def _contratti_e_documenti(self, blocco, rda, conto, centro) -> list[m.Contratto]:
        contratti: list[m.Contratto] = []
        corrente = None
        for riga in blocco:
            n = riga["_riga"]
            nuovo = self._contratto(riga, rda, conto, centro, n)
            if nuovo is not None:
                corrente = nuovo
                if nuovo not in contratti:
                    contratti.append(nuovo)
            if corrente is None:
                if valori(riga.get("fattura")) or valori(riga.get("ddt")):
                    self.anomalia(
                        n,
                        "DOCUMENTI_SENZA_CONTRATTO",
                        "fattura",
                        "",
                        "Fatture o DDT senza un contratto nella RDA: non importati.",
                    )
                continue
            fatture = self._fatture(riga, corrente, n)
            self._documenti(riga, corrente, fatture, rda, n)
        return contratti

    def _contratto(self, riga, rda, conto, centro, n) -> m.Contratto | None:
        sap = pulisci(riga.get("contratto"))
        codifica = pulisci(riga.get("codifica"))
        numero = re.sub(r"\s", "", sap or "")
        if not numero or "annullat" in numero.lower() or "chius" in numero.lower():
            numero = ""
        if not numero and codifica and RE_CONTRATTO.match(re.sub(r"\s", "", codifica)):
            numero = re.sub(r"\s", "", codifica)
        if not numero:
            return None
        if not RE_CONTRATTO.match(numero):
            self.anomalia(
                n,
                "NUMERO_CONTRATTO_NON_VALIDO",
                "contratto",
                numero,
                "Il numero di contratto SAP ha 10 cifre: riga senza contratto.",
            )
            return None
        esistente = m.Contratto.objects.filter(numero_contratto=numero, archiviato=False).first()
        if esistente is not None:
            if esistente.rda_id is None:
                esistente.rda = rda
                esistente._change_reason = "RDA d'origine dal Riepilogo"
                esistente.save(update_fields=["rda"])
                self.conta("contratti_collegati")
            elif esistente.rda_id != rda.pk:
                self.anomalia(
                    n,
                    "CONTRATTO_GIA_COLLEGATO",
                    "contratto",
                    numero,
                    "Il contratto ha già un'altra RDA d'origine: non modificato.",
                )
            return esistente
        return self._crea_contratto(riga, rda, numero, conto, centro, n)

    def _fornitore(self, nome, n) -> m.Fornitore | None:
        k = normalizza.chiave_nome(nome)
        if not k:
            return None
        f = self.fornitori.get(k)
        if f is None:
            f = m.Fornitore.objects.create(ragione_sociale=nome[:200])
            self.fornitori[k] = f
            self.conta("fornitori_creati")
            self.anomalia(
                n,
                "FORNITORE_CREATO",
                "fornitore",
                nome,
                "Fornitore non in anagrafica: creato, da verificare.",
            )
        return f

    def _fabbisogno(self, cluster: m.Cluster) -> m.Fabbisogno:
        descrizione = f"Acquisti una tantum {cluster.codice} (dal Riepilogo)"
        fab = m.Fabbisogno.objects.filter(
            descrizione=descrizione, cluster=cluster, archiviato=False
        ).first()
        if fab is None:
            fab = m.Fabbisogno.objects.create(
                descrizione=descrizione,
                cluster=cluster,
                categoria=self.categoria,
                note="Creato dall'import del Riepilogo RDA Ordini: contratti da ricollocare.",
            )
        return fab

    def _crea_contratto(self, riga, rda, numero, conto, centro, n) -> m.Contratto | None:
        nome = pulisci(riga.get("fornitore"))
        fornitore = self._fornitore(nome, n) if nome else None
        if fornitore is None:
            self.anomalia(
                n,
                "FORNITORE_MANCANTE",
                "fornitore",
                numero,
                "Contratto senza fornitore: non creato.",
            )
            return None
        emissione = data(riga.get("data_contratto"))
        inizio, fine = data(riga.get("inizio")), data(riga.get("fine"))
        una_tantum = not (inizio and fine)
        if una_tantum:
            inizio = emissione or rda.data_inserimento
            durata, fine_override = 1, None
        elif fine < inizio:
            self.anomalia(
                n,
                "DATE_INCOERENTI",
                "fine",
                f"{fine:%d/%m/%Y}",
                "Fine delle prestazioni prima dell'inizio: contratto una tantum.",
            )
            una_tantum, inizio, durata, fine_override = True, inizio, 1, None
        else:
            giorni = (fine - inizio).days + 1
            durata, fine_override = max(1, round(giorni / 30.4375)), fine
        base = importo(riga.get("importo_base")) or Decimal("0.00")
        massimo = importo(riga.get("importo_massimo"))
        contratto = m.Contratto(
            fabbisogno=self._fabbisogno(rda.cluster),
            fornitore=fornitore,
            descrizione=rda.descrizione,
            numero_contratto=numero,
            data_emissione=emissione,
            procedura=self.procedura,
            data_inizio=inizio,
            durata_mesi=durata,
            data_fine_override=fine_override,
            una_tantum=una_tantum,
            rda=rda,
            rup=rda.rup,
            cup=rda.cup,
            percorso_documenti=rda.percorso_file_server[: _lunghezza("percorso_documenti")],
            note="Creato dall'import del Riepilogo RDA Ordini.",
            chiave_import=f"riepilogo:{numero}",
        )
        contratto._change_reason = "Import del Riepilogo RDA Ordini"
        contratto.save()
        wbs = self._wbs(riga) or rda.posizioni.first().wbs
        voce = m.VoceContratto.objects.create(
            contratto=contratto,
            descrizione=rda.descrizione,
            conto=conto,
            codice_conto_importato=conto.codice,
            centro_costo=centro,
            natura=m.Natura.CAPEX if wbs else m.Natura.OPEX,
            wbs=wbs,
            importo_fisso=base,
            massimale_variabile=max(massimo - base, Decimal("0.00")) if massimo else Decimal(0),
        )
        m.PianoFatturazione.objects.create(voce=voce, modalita="da_verificare", periodicita="")
        self.conta("contratti_creati")
        self.anomalia(
            n,
            "CONTRATTO_CREATO",
            "contratto",
            numero,
            f"Contratto non registrato: creato nel fabbisogno «{contratto.fabbisogno}»"
            + (", una tantum." if una_tantum else "."),
        )
        return contratto

    def _chiudi(self, contratto: m.Contratto, n) -> None:
        if contratto.data_chiusura is not None:
            return
        contratto.data_chiusura = self.oggi
        contratto._change_reason = "Chiuso secondo il Riepilogo RDA Ordini"
        contratto.save(update_fields=["data_chiusura"])
        self.conta("contratti_chiusi")
        attestato = contratto.documenti_esecuzione.filter(
            archiviato=False, finale=True, tipo__in=["ATTESTAZIONE", "RICEZIONE"]
        ).exists()
        if not attestato:
            self.anomalia(
                n,
                "CHIUSO_SENZA_ATTESTAZIONE",
                "stato",
                contratto.numero_contratto or "",
                "«Chiuso» nel file, senza attestazione o modello di ricezione.",
            )

    # --- Fatture e documenti ----------------------------------------------------------------

    def _fatture(self, riga, contratto, n) -> list[m.Fattura]:
        numeri = valori(riga.get("fattura"))
        if not numeri:
            return []
        date = [data(v) for v in parti(riga.get("data_fattura"))] or [
            data(riga.get("data_fattura"))
        ]
        importi = [importo(v) for v in parti(riga.get("imponibile"))] or [None]
        if len(date) != len(numeri) or len(importi) != len(numeri):
            self.anomalia(
                n,
                "CELLE_DISALLINEATE",
                "fattura",
                "; ".join(numeri),
                "Numeri, date e importi delle fatture non si corrispondono: "
                "importate solo le coppie complete.",
            )
        trasferita = "trasferit" in (pulisci(riga.get("note")) or "").lower()
        create = []
        for i, numero in enumerate(numeri):
            d = date[i] if i < len(date) else None
            imp = importi[i] if i < len(importi) else None
            if d is None or imp is None:
                if len(date) == len(numeri) and len(importi) == len(numeri):
                    self.anomalia(
                        n,
                        "FATTURA_INCOMPLETA",
                        "fattura",
                        numero,
                        "Data o imponibile mancanti o illeggibili: non importata.",
                    )
                continue
            numero = numero[:60]
            if m.Fattura.objects.filter(
                fornitore=contratto.fornitore, numero=numero, archiviato=False
            ).exists():
                self.conta("fatture_gia_presenti")
                continue
            if contratto.data_chiusura is not None:
                self.anomalia(
                    n,
                    "FATTURA_DOPO_CHIUSURA",
                    "fattura",
                    numero,
                    "Fattura per un contratto chiuso: scartata.",
                )
                continue
            fattura = m.Fattura(
                contratto=contratto,
                fornitore=contratto.fornitore,
                numero=numero,
                data=d,
                importo_documento=imp,
                valore_base=imp,
                esito_convalida=m.Fattura.Esito.CONVALIDABILE,
                data_convalida=d,
                data_registrazione=self.oggi,
                data_trasferimento=d if trasferita else None,
                voce=contratto.voci.filter(archiviato=False).first(),
                note="Importata dal Riepilogo RDA Ordini.",
            )
            fattura._change_reason = "Import del Riepilogo RDA Ordini"
            fattura.save()
            create.append(fattura)
            self.conta("fatture_create")
        return create

    def _documento(self, contratto, tipo, numero, d, fattura=None, finale=False) -> None:
        numero = (numero or "")[:40]
        if m.DocumentoEsecuzione.objects.filter(
            contratto=contratto, tipo=tipo, numero=numero, archiviato=False
        ).exists():
            return
        doc = m.DocumentoEsecuzione(
            contratto=contratto,
            tipo=tipo,
            numero=numero,
            data=d,
            fattura=fattura,
            finale=finale,
            note="Importato dal Riepilogo RDA Ordini.",
        )
        doc._change_reason = "Import del Riepilogo RDA Ordini"
        doc.save()
        self.conta("documenti_creati")

    def _documenti(self, riga, contratto, fatture, rda, n) -> None:
        riserva = (
            (fatture[0].data if fatture else None)
            or contratto.data_emissione
            or (rda.data_inserimento)
        )
        fattura = fatture[0] if fatture else None
        Tipo = m.DocumentoEsecuzione.Tipo
        # DDT: numeri e date allineati per posizione.
        ddt = valori(riga.get("ddt"))
        date_ddt = [data(v) for v in parti(riga.get("data_ddt"))]
        for i, numero in enumerate(ddt):
            d = date_ddt[i] if i < len(date_ddt) and date_ddt[i] else None
            if d is None:
                self.anomalia(
                    n,
                    "DATA_NON_LEGGIBILE",
                    "ddt",
                    numero,
                    "DDT senza data leggibile: registrato con la data di riserva.",
                )
            self._documento(contratto, Tipo.DDT, numero, d or riserva)
        # Attestazioni e modelli di ricezione: righe con il tipo, date nella stessa cella o
        # nella colonna della data.
        date_att = [data(v) for v in parti(riga.get("data_attestazione"))]
        ultima_data = None
        k = 0
        for testo in valori(riga.get("attestazione")):
            if (d := data(testo)) is not None:
                ultima_data = d
                continue
            tipo = Tipo.RICEZIONE if "ricezion" in testo.lower() else Tipo.ATTESTAZIONE
            d = ultima_data or (date_att[k] if k < len(date_att) else None)
            k += 1
            if d is None:
                self.anomalia(
                    n,
                    "DATA_NON_LEGGIBILE",
                    "attestazione",
                    testo,
                    "Attestazione senza data: registrata con la data di riserva.",
                )
            self._documento(contratto, tipo, testo, d or riserva, finale=True)
            ultima_data = None
        # Numeri SAP: ODA/RDC, MAP, MIGO, legati alla prima fattura della riga.
        for campo, tipo in (("oda", Tipo.ODA), ("map", Tipo.MAP), ("migo", Tipo.MIGO)):
            for numero in valori(riga.get(campo)):
                self._documento(contratto, tipo, numero, riserva, fattura=fattura)


def _lunghezza(campo: str) -> int:
    return m.Contratto._meta.get_field(campo).max_length


def importa_riepilogo(
    percorso: str | Path, *, dry_run: bool = False, oggi: datetime.date | None = None
) -> Report:
    percorso = Path(percorso)
    oggi = oggi or timezone.localdate()
    report = Report(file=percorso.name, dry_run=dry_run)
    for chiave in (
        "rda_lette",
        "rda_create",
        "rda_gia_presenti",
        "fuori_perimetro",
        "non_importate",
        "contratti_collegati",
        "contratti_creati",
        "fornitori_creati",
        "fatture_create",
        "documenti_creati",
        "contratti_chiusi",
    ):
        report.conteggi[chiave] = 0
    foglio = leggi(percorso)
    try:
        with transaction.atomic():
            importatore = Importatore(report, oggi)
            for blocco in blocchi(foglio):
                with transaction.atomic():
                    importatore.importa_blocco(blocco)
            if dry_run:
                raise _Rollback
    except _Rollback:
        pass
    return report
