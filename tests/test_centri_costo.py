"""Centro di costo da tendina (tabella «Centri di costo», predefinito proposto). Dati di
fantasia."""

import datetime
from decimal import Decimal

import openpyxl
import pytest

from contratti import forms, forms_rda
from contratti import models as m
from contratti.importer.tabelle_sap import importa_tabelle_sap

from . import factories as f

pytestmark = pytest.mark.django_db


@pytest.fixture
def centri():
    m.CentroCosto.objects.create(codice="C0001", descrizione="Ufficio", predefinito=True)
    m.CentroCosto.objects.create(codice="C0002", descrizione="Altro")
    m.CentroCosto.objects.create(codice="C0009", descrizione="Dismesso", attivo=False)


def test_senza_tabella_resta_testo():
    form = forms_rda.PosizioneForm(prefix="p")
    assert form.fields["centro_costo"].__class__.__name__ == "CharField"


def test_posizione_nuova_propone_il_predefinito(centri):
    form = forms_rda.PosizioneForm(prefix="p")
    scelte = [v for v, _ in form.fields["centro_costo"].choices]
    assert scelte == ["", "C0001", "C0002"]  # solo gli attivi
    assert form.initial["centro_costo"] == "C0001"
    salvata = forms_rda.PosizioneForm(initial={"pk": 7, "centro_costo": "C0002"}, prefix="p")
    assert salvata.initial["centro_costo"] == "C0002"


def test_codice_registrato_fuori_tabella_resta(centri):
    form = forms_rda.PosizioneForm(initial={"pk": 7, "centro_costo": "XX99"}, prefix="p")
    assert ("XX99", "XX99 (non in tabella: sceglierne uno)") in form.fields["centro_costo"].choices
    dati = {"p-descrizione_breve": "Licenze", "p-quantita": "1", "p-importo_unitario": "10",
            "p-natura": "OPEX", "p-centro_costo": "C0009"}  # fmt: skip
    nuova = forms_rda.PosizioneForm(dati, prefix="p")
    assert not nuova.is_valid() and "centro_costo" in nuova.errors  # non attivo


def test_voce_del_contratto(centri):
    c = f.contratto()
    form = forms.VoceForm(contratto=c)
    assert form.initial["centro_costo"] == "C0001"
    voce = f.voce(c, centro_costo="XX99")
    form = forms.VoceForm(instance=voce, contratto=c)
    assert "XX99" in dict(form.fields["centro_costo"].choices)
    assert "centro_costo" not in form.initial or form.initial["centro_costo"] == "XX99"


def _xlsx(tmp_path, righe):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Centri di costo"
    for r in righe:
        ws.append(r)
    percorso = tmp_path / "centri.xlsx"
    wb.save(percorso)
    return percorso


def test_import_e_codici_usati_assenti(tmp_path):
    c = f.contratto()
    f.voce(c, centro_costo="C0001")
    f.voce(c, centro_costo="XX99")
    f.voce(c, centro_costo="XX99")
    report = importa_tabelle_sap(
        _xlsx(tmp_path, [["codice", "descrizione", "predefinito"], ["C0001", "Ufficio", "sì"]])
    )
    assert report.conteggi["centri_costo_creati"] == 1
    assert report.conteggi["centri_costo_non_in_tabella"] == 1
    [anomalia] = report.anomalie
    assert (anomalia.codice, anomalia.valore) == ("CENTRO_COSTO_NON_IN_TABELLA", "XX99")
    assert "2 voci" in anomalia.messaggio
    assert m.CentroCosto.codice_predefinito() == "C0001"


def test_import_esecuzione_segnala_il_centro_assente(centri, tmp_path):
    from contratti.importer.esecuzione import FOGLI, importa_esecuzione

    f.cluster("IT")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("RDA")
    ws.append(FOGLI["RDA"])
    riga = {"numero_rda": "5100000016", "cluster": "IT", "descrizione": "Prova",
            "posizione": "Hardware", "importo_unitario": 10, "centro_costo": "XX99"}  # fmt: skip
    ws.append([riga.get(c) for c in FOGLI["RDA"]])
    percorso = tmp_path / "esecuzione.xlsx"
    wb.save(percorso)
    report = importa_esecuzione(percorso, oggi=datetime.date(2026, 10, 9))
    assert [a.codice for a in report.anomalie] == ["CENTRO_COSTO_NON_IN_TABELLA"]
    posizione = m.PosizioneRDA.objects.get()
    assert posizione.centro_costo == "XX99" and posizione.importo_unitario == Decimal("10")
