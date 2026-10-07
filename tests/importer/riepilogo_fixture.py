"""Fixture anonimizzata del «Riepilogo RDA Ordini» (F11.4b): stessa struttura del file reale
(titoli, legenda, gruppi su celle unite, intestazione su due righe, una RDA su più righe unite,
celle con più valori, BOM, «'--», date seriali, conto e centro di costo invertiti, stati scritti
in modi diversi, righe fuori perimetro). Nomi, numeri e importi di fantasia."""

import datetime

import openpyxl

D = datetime.datetime

GRUPPI = [
    ("A", "I", "RICHIESTA D'ACQUISTO (RDA)"),
    ("J", "K", "RICHIESTA PROCESSATA DI ACQUISTO I DI CONTRATTO (RPA o RPC) - SOLO ANTE 2025"),
    ("L", "S", "CONTRATTO"),
    ("T", "V", "FATTURA"),
    ("W", "X", "DDT"),
    ("Y", "Z", "ATTESTAZIONE REGOLARE ESECUZIONE O MODELLO RICEZIONE MATERIALE"),
    ("AA", "AA", "PERCORSO NEL SERVER AZIENDALE"),
    ("AB", "AB", "STATO"),
    ("AC", "AC", "Basket riferimento masterplan Digitalizzazione"),
    ("AD", "AD", "NOTE"),
    ("AE", "AE", "SAP: ODA / RDC"),
    ("AF", "AF", "SAP: MAP"),
    ("AG", "AG", "SAP: MIGO"),
]  # fmt: skip
SOTTO = {
    "A": "Numero", "B": "Data/Data ultimo rilascio", "C": "Descrizione",
    "D": "Codice commessa/WBS", "E": "Centro di Costo", "F": "Codice Conto",
    "G": "Gestore/Richiedente SAP", "H": "RUP", "I": "CUP", "J": "Numero", "K": "Data",
    "L": "Numero (Codifica Siciliacque ante aprile 2025)", "M": "Numerazione SAP",
    "N": "Fornitore", "O": "Data", "P": "Data inizio prestazioni",
    "Q": "Data fine prestazione", "R": "Importo base", "S": "Importo massimo", "T": "Numero",
    "U": "Data", "V": "Imponibile", "W": "Numero", "X": "Data", "Y": "Numero", "Z": "Data",
}  # fmt: skip
PERCORSO = "smb://nas.esempio.local/Share/ICTDINN/00_IT/RDA_Ordini/2025/prova"

# Riga 8 in poi: (riga, {colonna: valore})
RIGHE = {
    # Prima di aprile 2025: fuori perimetro.
    8: {"A": " R2A 1/21", "B": D(2021, 1, 19), "C": "Acquisto di prova", "E": "C0001",
        "F": "G1280", "J": "RAA6/21", "L": "Richiesta chiusa", "M": "Richiesta chiusa",
        "AB": "Chiuso"},
    # RDA con tre contratti (righe 9-11) e una fattura sotto (riga 12); conto e centro di
    # costo invertiti; «Chiuso» senza attestazione.
    9: {"A": "5100000101", "B": D(2025, 5, 5), "C": "Fornitura di prova", "D": "--",
        "E": "600010", "F": "10520", "G": "999901 (Prova)", "H": "--", "I": "--",
        "J": "Non previsto in SAP", "K": "Non previsto in SAP", "L": "SI25000001",
        "M": "5800000101", "N": "Fornitore Prova B S.r.l.", "O": D(2025, 5, 12), "P": "--",
        "Q": "--", "R": 1000, "S": "--", "T": "F-1", "U": 45808, "V": 1000, "W": "'--",
        "X": "'--", "AA": PERCORSO, "AB": "Chiuso", "AD": "fattura trasferita",
        "AE": 3300000001},
    10: {"L": "SI25000002", "M": "5800000102", "N": "Fornitore Prova B S.r.l.",
         "O": D(2025, 6, 3), "R": 2308.6, "T": "F-2", "U": D(2025, 6, 30),
         "V": "\ufeff€ 2.308,60", "AD": "Fattura trasferita", "AE": 3300000002,
         "AF": "\n1000000001\n1000000002"},
    11: {"L": "SI25000003", "M": "\ufeff5800000103", "N": "Fornitore Prova B S.r.l.",
         "O": D(2025, 6, 12), "R": 500, "T": "'--", "U": "'--", "V": "'--"},
    12: {"T": "F-3", "U": D(2025, 7, 31), "V": 500, "W": "616/25", "X": 45804},
    # Contratto già registrato (5800000200): si collega; più fatture in una cella; DDT con
    # data illeggibile; modello di ricezione; stato «chiusa ».
    13: {"A": 5100000102, "B": D(2025, 6, 10), "C": "Servizio OT di prova",
         "D": "SI.99.PROVA", "E": "1OA699", "F": "G1517 600020", "G": "Prova", "H": "Rossi",
         "I": "\ufeffC00000000000001", "M": 5800000200, "N": "Fornitore Prova A",
         "O": D(2025, 6, 20), "P": D(2025, 7, 1), "Q": D(2026, 6, 30), "R": 3000,
         "T": "FA-10\n\nFA-11", "U": "01/07/2025\n15/07/2025", "V": "100\n200",
         "W": "DDT-1", "X": "24//11/25", "Y": "15/09/2025\nModello di ricezione materiale",
         "AB": "chiusa ", "AE": "3300000010\n3300000011", "AG": 5000000001},
    # Conto senza cluster: non è ICTDINN.
    14: {"A": "5100000103", "B": D(2025, 7, 1), "C": "Acquisto di un altro ufficio",
         "E": "C0002", "F": "700001", "M": "5800000400", "N": "Altro fornitore"},
    # Annullata.
    15: {"A": "5100000104", "B": D(2025, 8, 1), "C": "Richiesta annullata", "E": "C0001",
         "F": "600010", "L": "Annullato", "M": "Annullato"},
    # Numero tra apici, WBS, contratto con inizio e fine (non una tantum), importi testo.
    16: {"A": "'5100000105", "B": D(2025, 9, 1), "C": "Progetto di prova", "D": "WBS-TEST-01",
         "E": "C0001", "F": "600010", "M": "5800000300", "N": "Fornitore Prova A S.r.l.",
         "O": D(2025, 9, 10), "P": D(2025, 10, 1), "Q": D(2026, 9, 30), "R": "1.234,50",
         "S": 2000, "AB": "Aperto"},
    # Legenda in fondo.
    20: {"A": "Legenda"},
}  # fmt: skip
UNITE = [
    "A9:A11", "B9:B12", "C9:C12", "D9:D12", "E9:E12", "F9:F12", "G9:G12", "H9:H12",
    "I9:I12", "J9:J12", "K9:K12", "AA9:AA12", "AB9:AB11",
]  # fmt: skip


def crea(percorso):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RIEPILOGO"
    ws["A1"] = "Società di prova"
    ws["A2"] = "Area di prova"
    ws["A3"] = "RIEPILOGO RICHIESTE DI ACQUISTO, CONTRATTI, ATTESTATI, MODELLI DI RICEZIONE"
    ws["A4"] = "Aggiornamento del 02/10/26"
    ws["A5"], ws["B5"], ws["C5"] = "Legenda", "Rosso", "Campo non previsto in SAP"
    for inizio, fine, testo in GRUPPI:
        ws[f"{inizio}6"] = testo
        if inizio != fine:
            ws.merge_cells(f"{inizio}6:{fine}6")
    for colonna, testo in SOTTO.items():
        ws[f"{colonna}7"] = testo
    for riga, celle in RIGHE.items():
        for colonna, valore in celle.items():
            ws[f"{colonna}{riga}"] = valore
    for zona in UNITE:
        ws.merge_cells(zona)
    wb.save(percorso)
    return percorso
