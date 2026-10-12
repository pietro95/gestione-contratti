"""Lettura di file tabellari semplici (CSV o XLSX) per gli import di configurazione e dati."""

import csv
import datetime
from pathlib import Path

import openpyxl

from contratti.importer import normalizza

FORMATI_DATA = ("%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y")


class FileNonValido(Exception):
    """Il file non si può importare (colonne mancanti, valori illeggibili): nulla è scritto."""


def data(valore) -> datetime.date | None:
    """Data da cella Excel o testo; ``None`` se vuota; ``ValueError`` se illeggibile."""
    if isinstance(valore, datetime.datetime):
        return valore.date()
    if isinstance(valore, datetime.date):
        return valore
    testo = str(valore or "").strip()
    if not testo:
        return None
    for formato in FORMATI_DATA:
        try:
            return datetime.datetime.strptime(testo, formato).date()
        except ValueError:
            continue
    raise ValueError(testo)


def righe(percorso: str | Path, *, foglio: str | None = None) -> list[dict]:
    """Righe come dizionari con le intestazioni normalizzate (minuscole, senza spazi ai lati).

    XLSX: il foglio indicato se esiste, altrimenti il primo. CSV: separatore ; oppure ,
    (quello più presente nell'intestazione), codifica UTF-8 con o senza BOM.
    """
    percorso = Path(percorso)
    if percorso.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            wb = openpyxl.load_workbook(percorso, read_only=True, data_only=True)
        except Exception as exc:  # file danneggiato o non Excel
            raise FileNonValido(f"Il file non è un Excel leggibile ({exc}).") from exc
        ws = wb[foglio] if foglio and foglio in wb.sheetnames else wb.worksheets[0]
        tutte = list(ws.iter_rows(values_only=True))
        if not tutte:
            return []
        intestazione = [str(c or "").strip().lower() for c in tutte[0]]
        return [
            dict(zip(intestazione, r, strict=False))
            for r in tutte[1:]
            if any(v not in (None, "") for v in r)
        ]
    try:
        testo = percorso.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileNonValido("Il CSV non è in UTF-8: salvalo come «CSV UTF-8».") from exc
    linee = testo.splitlines()
    if not linee:
        return []
    separatore = ";" if linee[0].count(";") >= linee[0].count(",") else ","
    lettore = csv.DictReader(linee, delimiter=separatore)
    return [{(k or "").strip().lower(): v for k, v in r.items()} for r in lettore]


def colonna(tabella: list[dict], nome: str, alias: tuple[str, ...]) -> str:
    """Nome effettivo della colonna nel file, tra quelli ammessi; errore se manca."""
    presenti = set(tabella[0]) if tabella else set()
    for a in alias:
        if a in presenti:
            return a
    altri = f" (oppure {', '.join(alias[1:])})" if len(alias) > 1 else ""
    raise FileNonValido(f"Colonna mancante: {nome}{altri}.")


def codice(valore) -> str | None:
    """Codice da cella o testo: 5800001234.0 di Excel diventa 5800001234; nulli → None."""
    if isinstance(valore, float) and valore.is_integer():
        valore = int(valore)
    return normalizza.codice(valore)


MESSAGGIO_CENTRO_COSTO = (
    "Centro di costo assente dalla tabella: si registra così e si sistema dalla tendina."
)


def centro_costo_da_segnalare(codice: str | None) -> bool:
    """Il codice non è nella tabella dei centri di costo (solo se la tabella è caricata)."""
    from contratti import models as m

    if not codice or not m.CentroCosto.objects.exists():
        return False
    return not m.CentroCosto.objects.filter(codice=codice).exists()
