"""Lettura degli Excel attuali con openpyxl e normalizzazione riga per riga (pura)."""

import datetime
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import openpyxl

from contratti.importer import normalizza as n
from contratti.importer.normalizza import Anomalia

# Intestazioni dei file, con le varianti trovate tra IT e TLC.
COLONNE = {
    "contratto": ("contratto",),
    "cig": ("cig",),
    "conto_nuovo": ("nuovo codice conto da 01/01/26", "codice conto dal 01/01/26"),
    "conto": ("codice conto",),
    "wbs": ("codice wbs",),
    "descrizione": ("descrizione",),
    "fornitore": ("fornitore",),
    "minimo": ("corrispettivo minimo",),
    "extra": ("corrispettivo extra",),
    "massimo": ("corrispettivo massimo",),
    "data": ("data di attivazione",),
    "mesi": ("mesi contratto",),
    "stato": ("stato",),
    "note": ("note",),
}

# Mappatura degli stati (SPEC §7.7).
DECISIONE_DA_STATO = {
    "non rinnovare": "dismettere",
    "rinnovato": "rinnovare",
    "attivo": "rinnovare",
}
PRATICA_DA_STATO = {
    "offerta richiesta": "offerta_richiesta",
    "offerta ricevuta": "offerta_ricevuta",
    "rda richiesta": "rda_richiesta",
    "rda predisposta": "rda_predisposta",
}
# Si ricavano dalle date: nessun dato da importare.
STATI_DERIVATI = frozenset({"scaduto", "in scadenza", "da rinnovare"})


@dataclass(frozen=True)
class RigaGrezza:
    foglio: str
    numero: int  # numero di riga nel foglio Excel
    valori: dict
    massimo_formula: bool  # False = valore scritto a mano al posto della formula


@dataclass
class Riga:
    """Riga normalizzata, pronta per il caricamento."""

    foglio: str
    numero: int
    finanziato: bool
    numero_oda: str | None
    cig: str | None
    conto: str | None
    conto_vecchio: str | None
    wbs: str | None
    descrizione: str | None
    fornitore: str | None
    minimo: Decimal | None
    extra: Decimal | None
    data_inizio: datetime.date | None
    mesi: int | None
    decisione: str
    stato_pratica: str | None
    numero_rda: str | None
    note: str
    anomalie: list[Anomalia] = field(default_factory=list)


@dataclass
class Libro:
    righe: list[RigaGrezza]
    fornitori: list[tuple[str, str | None]]  # (nome, codice SAP) dal foglio Fornitori
    anomalie_fogli: list[tuple[str, Anomalia]] = field(default_factory=list)


def _intestazioni(riga) -> dict[str, int] | None:
    nomi = [n.testo(c.value) for c in riga]
    if not nomi or (nomi[0] or "").lower() != "contratto":
        return None
    indici = {}
    for chiave, varianti in COLONNE.items():
        for i, nome in enumerate(nomi):
            if nome and nome.lower() in varianti:
                indici[chiave] = i
                break
    return indici


def leggi(percorso: str | Path) -> Libro:
    """Fogli "Contratti…" e "Fornitori". Righe vuote e riga "Totale" escluse."""
    wb = openpyxl.load_workbook(percorso, data_only=False)
    righe: list[RigaGrezza] = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith("contratti"):
            righe += _leggi_foglio_contratti(ws)
    fornitori = []
    if "Fornitori" in wb.sheetnames:
        for r in list(wb["Fornitori"].iter_rows(values_only=True))[1:]:
            nome = n.testo(r[0]) if r else None
            if nome:
                fornitori.append((nome, n.codice(r[1]) if len(r) > 1 else None))
    return Libro(righe=righe, fornitori=fornitori)


def _leggi_foglio_contratti(ws) -> list[RigaGrezza]:
    indici = None
    risultato = []
    for riga in ws.iter_rows():
        if indici is None:
            indici = _intestazioni(riga)
            continue
        valori = {k: riga[i].value if i < len(riga) else None for k, i in indici.items()}
        primo = n.testo(riga[0].value)
        if primo and primo.lower() == "totale":
            break
        utili = {k: v for k, v in valori.items() if k not in ("stato", "note", "massimo")}
        if all(n.testo(v) is None for v in utili.values()):
            continue
        massimo = valori.get("massimo")
        formula = isinstance(massimo, str) and massimo.startswith("=")
        risultato.append(
            RigaGrezza(ws.title, riga[0].row, valori, massimo_formula=formula or massimo is None)
        )
    return risultato


def finanziato(foglio: str) -> bool:
    """I fogli dei contratti finanziati e dei dissalatori: spese su altre WBS."""
    nome = foglio.lower()
    return "finanziat" in nome or "dissalator" in nome


def normalizza_riga(g: RigaGrezza) -> Riga:
    v = g.valori
    anomalie: list[Anomalia] = []

    for campo in ("contratto", "cig"):
        if n.aveva_invisibili(v.get(campo)):
            anomalie.append(
                Anomalia("CARATTERI_INVISIBILI", campo, repr(v.get(campo)), "Rimossi BOM/spazi.")
            )
    numero_oda, a = n.oda(v.get("contratto"))
    anomalie += a
    cig, a = n.cig(v.get("cig"))
    anomalie += a

    minimo, a = n.importo(v.get("minimo"), "minimo")
    anomalie += a
    extra, a = n.importo(v.get("extra"), "extra")
    anomalie += a
    if not g.massimo_formula:
        massimo, _ = n.importo(v.get("massimo"), "massimo")
        anomalie.append(
            Anomalia(
                "MASSIMO_SOVRASCRITTO",
                "massimo",
                str(massimo),
                "Corrispettivo Massimo scritto a mano al posto della formula.",
            )
        )
        atteso = (minimo or Decimal(0)) + (extra or Decimal(0))
        if massimo is not None and massimo != atteso:
            anomalie.append(
                Anomalia(
                    "MASSIMO_INCOERENTE",
                    "massimo",
                    str(massimo),
                    f"Massimo {massimo} diverso da minimo + extra = {atteso}.",
                )
            )
    if (
        (minimo or 0) == 0
        and (extra or 0) == 0
        and not any(x.codice == "IMPORTO_NON_LEGGIBILE" for x in anomalie)
    ):
        anomalie.append(Anomalia("IMPORTO_ZERO", "minimo", "0", "Contratto con valore zero."))

    data_inizio = n.data(v.get("data"))
    mesi = n.intero(v.get("mesi"))
    if data_inizio is None:
        anomalie.append(
            Anomalia(
                "DATA_MANCANTE",
                "data",
                str(v.get("data")),
                "Data di attivazione mancante: escluso dai calcoli di budget.",
            )
        )
    if not mesi:
        anomalie.append(
            Anomalia("DURATA_MANCANTE", "mesi", str(v.get("mesi")), "Durata assunta di 12 mesi.")
        )

    descrizione = n.testo(v.get("descrizione"))
    if descrizione is None:
        anomalie.append(Anomalia("DESCRIZIONE_MANCANTE", "descrizione", "", "Descrizione vuota."))
    fornitore = n.testo(v.get("fornitore"))
    if fornitore is None:
        anomalie.append(Anomalia("FORNITORE_MANCANTE", "fornitore", "", "Fornitore non indicato."))

    fin = finanziato(g.foglio)
    wbs = n.codice(v.get("wbs"))
    if fin and wbs is None:
        anomalie.append(
            Anomalia("WBS_MANCANTE", "wbs", "", "Contratto finanziato senza WBS: da definire.")
        )

    stato = (n.testo(v.get("stato")) or "").lower()
    decisione = DECISIONE_DA_STATO.get(stato, "rinnovare" if stato in PRATICA_DA_STATO else "")
    if not decisione:
        decisione = "da_valutare"
    pratica = PRATICA_DA_STATO.get(stato)
    if (
        stato
        and stato not in DECISIONE_DA_STATO | PRATICA_DA_STATO
        and stato not in (STATI_DERIVATI)
    ):
        anomalie.append(Anomalia("STATO_SCONOSCIUTO", "stato", stato, "Stato non riconosciuto."))

    note = n.testo(v.get("note")) or ""
    return Riga(
        foglio=g.foglio,
        numero=g.numero,
        finanziato=fin,
        numero_oda=numero_oda,
        cig=cig,
        conto=n.codice(v.get("conto_nuovo")) or n.codice(v.get("conto")),
        conto_vecchio=n.codice(v.get("conto")),
        wbs=wbs,
        descrizione=descrizione,
        fornitore=fornitore,
        minimo=minimo,
        extra=extra,
        data_inizio=data_inizio,
        mesi=mesi or 12,
        decisione=decisione,
        stato_pratica=pratica,
        numero_rda=n.rda_da_nota(note),
        note=note,
        anomalie=anomalie,
    )
