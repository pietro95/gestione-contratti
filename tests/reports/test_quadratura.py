"""Quadratura con il Dashboard Excel (SPEC §7.8): formula Excel ricalcolata, bug compreso."""

import datetime
from decimal import Decimal

import openpyxl
import pytest
from django.core.management import call_command

from contratti.engine.tipi import Livello
from contratti.importer.carica import importa
from reports.budget import Filtri, calcola_budget
from reports.quadratura import excel_anno, mesi_toccati, quadra, scadenza_excel

from ..importer.conftest import IT, excel, riga  # noqa: F401 - fixture

pytestmark = pytest.mark.django_db
D = datetime.date
ANNI = [2024, 2025, 2026, 2027, 2028]


def test_mesi_toccati_e_bug_dei_mesi():
    # 12 mesi dal 21/12/2024: l'Excel tocca 13 mesi di calendario (SPEC §2.1).
    fine, _ = scadenza_excel("=EDATE(...)-1", D(2024, 12, 21), 12)
    assert fine == D(2025, 12, 20)
    assert mesi_toccati(D(2024, 12, 21), fine) == 13
    totale = sum(excel_anno(Decimal(1200), 12, D(2024, 12, 21), fine, a) for a in (2024, 2025))
    assert totale == Decimal(1300)  # 108 %
    assert excel_anno(Decimal(1200), 12, D(2024, 12, 21), fine, 2026) == 0


def test_scadenza_scritta_a_mano():
    assert scadenza_excel(datetime.datetime(2026, 6, 30), D(2025, 1, 1), 12) == (
        D(2026, 6, 30),
        True,
    )


def test_quadratura_riga_per_riga(excel, conti):  # noqa: F811
    percorso = excel(
        [
            riga(
                **{
                    "Data di attivazione": datetime.datetime(2024, 12, 21),
                    "Corrispettivo Minimo": 1200.0,
                }
            ),
            riga(
                Contratto="5800000002",
                Cig="Z000000002",
                Descrizione="Canone allineato",
                **{
                    "Data di attivazione": datetime.datetime(2025, 1, 1),
                    "Corrispettivo Minimo": 3650.0,
                },
            ),
            riga(
                Contratto="5800000003",
                Cig="Z000000003",
                Descrizione="Massimo a mano",
                **{"Corrispettivo Extra": 100.0, "Corrispettivo Massimo": 999.0},
            ),
        ]
    )
    prima = quadra(percorso, area="IT", anni=ANNI)
    assert all(r.cause[0] == "NON_IMPORTATA" for r in prima.righe)

    importa(percorso, area="IT")
    q = quadra(percorso, area="IT", anni=ANNI)
    bug, allineata, massimo = q.righe
    assert "BUG_MESI" in bug.cause
    assert bug.app_fisso[2024] + bug.app_fisso[2025] == Decimal("1200.00")
    assert allineata.cause == []  # anno solare intero: mesi e giorni coincidono
    assert allineata.differenza(2025) == 0
    assert "MASSIMO_SOVRASCRITTO" in massimo.cause
    assert massimo.excel_massimo[2026] == Decimal(999)
    riepilogo = {r["anno"]: r for r in q.riepilogo()}
    assert (
        riepilogo[2025]["diff_minimo"]
        == riepilogo[2025]["app_fisso"] - riepilogo[2025]["excel_minimo"]
    )


def test_app_coincide_con_il_budget_consolidato(conti):
    importa(IT, area="IT")
    q = quadra(IT, area="IT", anni=[2026])
    assert not q.escluse or all("mancant" in motivo.lower() for _, motivo in q.escluse)
    assert {c for r in q.righe for c in r.cause} >= {"BUG_MESI"}
    assert not any("NON_IMPORTATA" in r.cause for r in q.righe)
    # Somma delle righe del foglio Contratti ≤ competenza dell'area (c'è anche il foglio
    # dei dissalatori, fuori dal Dashboard).
    budget = calcola_budget(Filtri(anno=2026, area="IT", previsionali=False), D(2026, 10, 1))
    competenza = budget.totali()["consolidato"]["competenza"][Livello.FISSO]
    assert sum(r.app_fisso[2026] for r in q.righe) <= competenza


def test_comando_scrive_il_report(tmp_path, conti):
    importa(IT, area="IT")
    destinazione = tmp_path / "quadratura.xlsx"
    call_command(
        "quadratura_excel",
        str(IT),
        "--area",
        "IT",
        "--anni",
        "2025-2027",
        "--report",
        str(destinazione),
    )
    wb = openpyxl.load_workbook(destinazione)
    assert wb.sheetnames == ["Riepilogo", "Dettaglio righe", "Cause ed esclusioni"]
    anni = [r[0] for r in wb["Riepilogo"].iter_rows(min_row=6, values_only=True)]
    assert anni == [2025, 2026, 2027]
